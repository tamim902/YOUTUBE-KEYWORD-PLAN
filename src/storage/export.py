"""Spreadsheet export service supporting CSV and XLSX formats.

Features:
- Protection against CSV and spreadsheet formula injection (=, +, -, @, \\t, \\r).
- UTF-8 with Byte Order Mark (UTF-8 BOM) for flawless Bangla and multi-script display in Excel.
- Clean openpyxl XLSX generation with styled header rows and autowidth columns.
- Metadata and summary headers.
- Support for project, selected, and filtered keyword scopes.
- Safe empty-result handling.
- Descriptive filename generation.
"""

import io
import json
import re
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Characters that trigger spreadsheet formula execution in Microsoft Excel, LibreOffice Calc, and Google Sheets
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def sanitize_formula_injection(value: Any) -> Any:
    """Neutralize spreadsheet formula injection by prefixing formula indicators with a single quote.

    Leaves non-string values or safe strings untouched.
    Ensures spreadsheet applications treat untrusted cell values strictly as plain text.
    Preserves original data in the database while protecting the user at the export boundary.
    """
    if isinstance(value, str) and value:
        # Check if the string begins with any known formula injection trigger
        if value.startswith(FORMULA_PREFIXES):
            return f"'{value}"
    return value


def sanitize_dataframe_for_export(df: pd.DataFrame) -> pd.DataFrame:
    """Create a sanitized copy of a DataFrame with formula triggers escaped."""
    if df.empty:
        return df.copy()

    clean_df = df.copy()
    for col in clean_df.columns:
        if clean_df[col].dtype == object or clean_df[col].dtype == "string":
            clean_df[col] = clean_df[col].apply(sanitize_formula_injection)
    return clean_df


def generate_export_filename(
    base_name: str,
    scope: str = "all",
    extension: str = "csv",
) -> str:
    """Generate a clean, sanitized, and descriptive export filename.

    Example: youtube_keywords_python_tutorial_all_20261018_203000.csv
    """
    clean_base = re.sub(r"[^\w\s-]", "", base_name).strip().replace(" ", "_")
    if not clean_base:
        clean_base = "youtube_keywords"
    else:
        clean_base = f"youtube_keywords_{clean_base}"

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    clean_ext = extension.lstrip(".")
    return f"{clean_base}_{scope}_{timestamp}.{clean_ext}"


class ExportService:
    """Orchestrates CSV and Excel exports with formula protection and UTF-8 encoding."""

    @staticmethod
    def export_csv(
        df: pd.DataFrame,
        sanitize_formulas: bool = True,
        include_bom: bool = True,
        metadata_header: Optional[Dict[str, Any]] = None,
    ) -> bytes:
        """Serialize a DataFrame to CSV bytes with UTF-8 BOM encoding and formula sanitization.

        Why UTF-8 BOM (utf-8-sig)?
        Microsoft Excel on Windows defaults to the legacy Windows ANSI code page unless
        a Byte Order Mark (BOM: \\xef\\xbb\\xbf) is present. Using utf-8-sig ensures Excel
        renders non-Latin scripts (such as Bengali বাংলা) with 100% typography fidelity without mojibake.
        """
        export_df = sanitize_dataframe_for_export(df) if sanitize_formulas else df.copy()

        buf = io.StringIO()

        # Optional metadata comments at top of file
        if metadata_header:
            buf.write("# YouTube Keyword Finder Export Metadata\n")
            for k, v in metadata_header.items():
                buf.write(f"# {k}: {v}\n")
            buf.write("#\n")

        export_df.to_csv(buf, index=False)
        csv_text = buf.getvalue()

        encoding = "utf-8-sig" if include_bom else "utf-8"
        return csv_text.encode(encoding)

    @staticmethod
    def export_excel(
        df: pd.DataFrame,
        sanitize_formulas: bool = True,
        sheet_name: str = "Keywords",
        metadata_header: Optional[Dict[str, Any]] = None,
    ) -> bytes:
        """Serialize a DataFrame to styled openpyxl XLSX bytes with formula sanitization."""
        export_df = sanitize_dataframe_for_export(df) if sanitize_formulas else df.copy()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name[:31]  # Excel limits sheet names to 31 chars

        # Styling definitions
        header_fill = PatternFill(
            start_color="1F2937", end_color="1F2937", fill_type="solid"
        )
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        meta_font = Font(name="Calibri", size=9, italic=True, color="6B7280")
        cell_font = Font(name="Calibri", size=10)
        thin_border = Border(
            left=Side(style="thin", color="E5E7EB"),
            right=Side(style="thin", color="E5E7EB"),
            top=Side(style="thin", color="E5E7EB"),
            bottom=Side(style="thin", color="E5E7EB"),
        )

        current_row = 1

        # Optional metadata block
        if metadata_header:
            ws.cell(
                row=current_row,
                column=1,
                value="YouTube Keyword Finder — Research Export Metadata",
            ).font = Font(name="Calibri", size=11, bold=True, color="111827")
            current_row += 1
            for k, v in metadata_header.items():
                cell = ws.cell(row=current_row, column=1, value=f"{k}: {v}")
                cell.font = meta_font
                current_row += 1
            current_row += 1  # Blank row before table

        # Write table headers
        headers = list(export_df.columns)
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=str(header))
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[current_row].height = 24
        header_row_idx = current_row
        current_row += 1

        # Write rows
        if not export_df.empty:
            for row_tuple in export_df.itertuples(index=False):
                for col_idx, val in enumerate(row_tuple, start=1):
                    # Represent lists or dicts nicely if any exist
                    cell_val = val
                    if isinstance(val, (list, tuple)):
                        cell_val = ", ".join(str(x) for x in val)
                    cell = ws.cell(row=current_row, column=col_idx, value=cell_val)
                    cell.font = cell_font
                    cell.border = thin_border
                    if isinstance(cell_val, (int, float)):
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                ws.row_dimensions[current_row].height = 19
                current_row += 1

        # Adjust column widths automatically
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # Only check table rows (skip top metadata rows to avoid skewing width)
                if cell.row >= header_row_idx and cell.value:
                    val_str = str(cell.value)
                    max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Freeze header row
        ws.freeze_panes = ws.cell(row=header_row_idx + 1, column=1)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    @staticmethod
    def export_serp_csv(analysis: Any) -> bytes:
        """Export KeywordSerpAnalysis videos and metrics to CSV bytes with formula protection."""
        meta = {
            "Keyword": getattr(analysis, "keyword", ""),
            "Language": getattr(analysis, "language", "en"),
            "Country": getattr(analysis, "country", "US"),
            "Opportunity Score": f"{getattr(analysis, 'opportunity_score', 0.0)} / 100 ({getattr(analysis, 'opportunity_label', '')})",
            "Average Views": f"{int(getattr(analysis, 'average_views', 0)):,}",
            "Median Views": f"{int(getattr(analysis, 'median_views', 0)):,}",
            "Total Videos Analyzed": getattr(analysis, "total_videos_analyzed", 0),
            "Fetched At": getattr(analysis, "fetched_at", ""),
        }
        rows = []
        for v in getattr(analysis, "videos", []):
            rows.append(
                {
                    "Rank": getattr(v, "rank", 1),
                    "Title": getattr(v, "title", ""),
                    "Channel": getattr(v, "channel_title", ""),
                    "Views": getattr(v, "view_count", 0),
                    "Likes": getattr(v, "like_count", 0) if getattr(v, "like_count", None) is not None else 0,
                    "Comments": getattr(v, "comment_count", 0) if getattr(v, "comment_count", None) is not None else 0,
                    "Duration": getattr(v, "duration", ""),
                    "Published Date": getattr(v, "publish_date", ""),
                    "Age": getattr(v, "relative_age", ""),
                    "URL": getattr(v, "video_url", ""),
                    "Tags": ", ".join(getattr(v, "tags", [])) if getattr(v, "tags", None) else "",
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_csv(
            df, sanitize_formulas=True, include_bom=True, metadata_header=meta
        )

    @staticmethod
    def export_serp_excel(analysis: Any) -> bytes:
        """Export KeywordSerpAnalysis videos and metrics to styled XLSX bytes with formula protection."""
        meta = {
            "Keyword": getattr(analysis, "keyword", ""),
            "Language": getattr(analysis, "language", "en"),
            "Country": getattr(analysis, "country", "US"),
            "Opportunity Score": f"{getattr(analysis, 'opportunity_score', 0.0)} / 100 ({getattr(analysis, 'opportunity_label', '')})",
            "Average Views": f"{int(getattr(analysis, 'average_views', 0)):,}",
            "Median Views": f"{int(getattr(analysis, 'median_views', 0)):,}",
            "Total Videos Analyzed": getattr(analysis, "total_videos_analyzed", 0),
            "Fetched At": getattr(analysis, "fetched_at", ""),
        }
        rows = []
        for v in getattr(analysis, "videos", []):
            rows.append(
                {
                    "Rank": getattr(v, "rank", 1),
                    "Title": getattr(v, "title", ""),
                    "Channel": getattr(v, "channel_title", ""),
                    "Views": getattr(v, "view_count", 0),
                    "Likes": getattr(v, "like_count", 0) if getattr(v, "like_count", None) is not None else 0,
                    "Comments": getattr(v, "comment_count", 0) if getattr(v, "comment_count", None) is not None else 0,
                    "Duration": getattr(v, "duration", ""),
                    "Published Date": getattr(v, "publish_date", ""),
                    "Age": getattr(v, "relative_age", ""),
                    "URL": getattr(v, "video_url", ""),
                    "Tags": ", ".join(getattr(v, "tags", [])) if getattr(v, "tags", None) else "",
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_excel(
            df,
            sanitize_formulas=True,
            sheet_name="SERP Competitors",
            metadata_header=meta,
        )

    @staticmethod
    def export_opportunities_csv(opportunities: List[Any]) -> bytes:
        """Export content opportunities to CSV bytes with formula protection and UTF-8 BOM."""
        rows = []
        for opp in opportunities:
            d = opp.to_dict() if hasattr(opp, "to_dict") else dict(opp)
            sup = d.get("supporting_keywords", [])
            sup_str = ", ".join(sup) if isinstance(sup, list) else str(sup)
            rows.append(
                {
                    "Topic": d.get("topic", ""),
                    "Primary Keyword": d.get("primary_keyword", ""),
                    "Supporting Keywords": sup_str,
                    "Target Audience": d.get("target_audience", "General Audience"),
                    "Search Intent": d.get("search_intent", "Informational"),
                    "Format": d.get("content_format", "Step-by-step Guide"),
                    "Opportunity Score": d.get("opportunity_score", 50.0),
                    "Confidence Level": d.get("confidence_level", "Moderate Evidence"),
                    "Status": d.get("status", "Idea"),
                    "Priority": d.get("priority", "Medium"),
                    "Created At": d.get("created_at", ""),
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_csv(df, sanitize_formulas=True, include_bom=True)

    @staticmethod
    def export_opportunities_excel(opportunities: List[Any]) -> bytes:
        """Export content opportunities to styled XLSX bytes with formula protection."""
        rows = []
        for opp in opportunities:
            d = opp.to_dict() if hasattr(opp, "to_dict") else dict(opp)
            sup = d.get("supporting_keywords", [])
            sup_str = ", ".join(sup) if isinstance(sup, list) else str(sup)
            rows.append(
                {
                    "Topic": d.get("topic", ""),
                    "Primary Keyword": d.get("primary_keyword", ""),
                    "Supporting Keywords": sup_str,
                    "Target Audience": d.get("target_audience", "General Audience"),
                    "Search Intent": d.get("search_intent", "Informational"),
                    "Format": d.get("content_format", "Step-by-step Guide"),
                    "Opportunity Score": d.get("opportunity_score", 50.0),
                    "Confidence Level": d.get("confidence_level", "Moderate Evidence"),
                    "Status": d.get("status", "Idea"),
                    "Priority": d.get("priority", "Medium"),
                    "Created At": d.get("created_at", ""),
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_excel(df, sanitize_formulas=True, sheet_name="Opportunities")

    @staticmethod
    def export_briefs_csv(briefs: List[Any]) -> bytes:
        """Export content briefs to CSV bytes with formula protection."""
        rows = []
        for b in briefs:
            d = b.to_dict() if hasattr(b, "to_dict") else dict(b)
            rows.append(
                {
                    "Working Title": d.get("working_title", ""),
                    "Primary Keyword": d.get("primary_keyword", ""),
                    "Target Audience": d.get("target_audience", ""),
                    "Search Intent": d.get("search_intent", ""),
                    "Viewer Problem": d.get("viewer_problem", ""),
                    "Unique Angle": d.get("unique_angle", ""),
                    "Proposed Intro": d.get("proposed_intro", ""),
                    "Main Sections": " | ".join(d.get("main_sections", [])) if isinstance(d.get("main_sections"), list) else str(d.get("main_sections", "")),
                    "Conclusion & CTA": d.get("conclusion_cta", ""),
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_csv(df, sanitize_formulas=True, include_bom=True)

    @staticmethod
    def export_briefs_excel(briefs: List[Any]) -> bytes:
        """Export content briefs to styled XLSX bytes with formula protection."""
        rows = []
        for b in briefs:
            d = b.to_dict() if hasattr(b, "to_dict") else dict(b)
            rows.append(
                {
                    "Working Title": d.get("working_title", ""),
                    "Primary Keyword": d.get("primary_keyword", ""),
                    "Target Audience": d.get("target_audience", ""),
                    "Search Intent": d.get("search_intent", ""),
                    "Viewer Problem": d.get("viewer_problem", ""),
                    "Unique Angle": d.get("unique_angle", ""),
                    "Proposed Intro": d.get("proposed_intro", ""),
                    "Main Sections": " | ".join(d.get("main_sections", [])) if isinstance(d.get("main_sections"), list) else str(d.get("main_sections", "")),
                    "Conclusion & CTA": d.get("conclusion_cta", ""),
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_excel(df, sanitize_formulas=True, sheet_name="Content Briefs")

    @staticmethod
    def export_calendar_csv(entries: List[Any]) -> bytes:
        """Export content calendar entries to CSV bytes with formula protection."""
        rows = []
        for e in entries:
            d = e.to_dict() if hasattr(e, "to_dict") else dict(e)
            rows.append(
                {
                    "Planned Date": d.get("planned_date", ""),
                    "Title": d.get("title", ""),
                    "Primary Keyword": d.get("primary_keyword", ""),
                    "Status": d.get("status", "Idea"),
                    "Priority": d.get("priority", "Medium"),
                    "Notes": d.get("notes", ""),
                    "Created At": d.get("created_at", ""),
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_csv(df, sanitize_formulas=True, include_bom=True)

    @staticmethod
    def export_calendar_excel(entries: List[Any]) -> bytes:
        """Export content calendar entries to styled XLSX bytes."""
        rows = []
        for e in entries:
            d = e.to_dict() if hasattr(e, "to_dict") else dict(e)
            rows.append(
                {
                    "Planned Date": d.get("planned_date", ""),
                    "Title": d.get("title", ""),
                    "Primary Keyword": d.get("primary_keyword", ""),
                    "Status": d.get("status", "Idea"),
                    "Priority": d.get("priority", "Medium"),
                    "Notes": d.get("notes", ""),
                    "Created At": d.get("created_at", ""),
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_excel(df, sanitize_formulas=True, sheet_name="Content Calendar")

    @staticmethod
    def export_overlap_report_csv(overlaps: List[Any]) -> bytes:
        """Export topic overlap & cannibalization results to CSV bytes."""
        rows = []
        for ov in overlaps:
            d = ov.to_dict() if hasattr(ov, "to_dict") else dict(ov)
            shared = d.get("shared_keywords", [])
            shared_str = ", ".join(shared) if isinstance(shared, list) else str(shared)
            rows.append(
                {
                    "Topic A": d.get("topic_a", ""),
                    "Topic B": d.get("topic_b", ""),
                    "Similarity Score": d.get("similarity_score", 0.0),
                    "Overlap Type": d.get("overlap_type", ""),
                    "Shared Keywords": shared_str,
                    "Warning Message": d.get("warning_message", ""),
                    "Resolution": d.get("resolution", "Unresolved"),
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_csv(df, sanitize_formulas=True, include_bom=True)

    @staticmethod
    def export_overlap_report_excel(overlaps: List[Any]) -> bytes:
        """Export topic overlap & cannibalization results to styled XLSX bytes."""
        rows = []
        for ov in overlaps:
            d = ov.to_dict() if hasattr(ov, "to_dict") else dict(ov)
            shared = d.get("shared_keywords", [])
            shared_str = ", ".join(shared) if isinstance(shared, list) else str(shared)
            rows.append(
                {
                    "Topic A": d.get("topic_a", ""),
                    "Topic B": d.get("topic_b", ""),
                    "Similarity Score": d.get("similarity_score", 0.0),
                    "Overlap Type": d.get("overlap_type", ""),
                    "Shared Keywords": shared_str,
                    "Warning Message": d.get("warning_message", ""),
                    "Resolution": d.get("resolution", "Unresolved"),
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_excel(df, sanitize_formulas=True, sheet_name="Cannibalization Report")

    @staticmethod
    def export_products_csv(products: List[Any], session_info: Optional[Dict[str, Any]] = None) -> bytes:
        """Export discovered products to CSV bytes with formula sanitization and UTF-8 BOM."""
        rows = []
        for p in products:
            d = p.to_dict() if hasattr(p, "to_dict") else dict(p)
            bd = d.get("opportunity_score_breakdown", {}) or {}
            q_list = d.get("queries_executed", [])
            q_str = "; ".join(q_list) if isinstance(q_list, list) else str(q_list)
            audit = d.get("audit_trail", []) or []
            rows.append(
                {
                    "Product Name": d.get("product_name", ""),
                    "Category": d.get("category", ""),
                    "Brand": d.get("brand", ""),
                    "Model": d.get("model", ""),
                    "Priority": d.get("priority", "Medium"),
                    "Niche": d.get("niche", ""),
                    "Classification": d.get("classification", ""),
                    "Classification Reason": d.get("classification_reason", ""),
                    "Matching Videos Count": d.get("matching_videos_count", 0),
                    "Unique Channels Count": d.get("unique_channels_count", 0),
                    "Latest Video Date": d.get("latest_video_date", "") or "None",
                    "Opportunity Score": d.get("opportunity_score", 50.0),
                    "Confidence Level": d.get("confidence_level", "Low"),
                    "Scarcity Points (0-40)": bd.get("opportunity_factor", 0.0),
                    "Coverage Points (0-25)": bd.get("search_coverage_factor", 0.0),
                    "Identity Points (0-20)": bd.get("identity_confidence_factor", 0.0),
                    "Relevance Points (0-15)": bd.get("niche_relevance_factor", 0.0),
                    "Source": d.get("source", "Extracted Candidate"),
                    "Evidence URL": d.get("evidence_url", ""),
                    "Evidence Reference": d.get("evidence_reference", ""),
                    "Verification Status": d.get("verification_status", "Extracted Candidate"),
                    "Original Name": d.get("original_name", ""),
                    "User Corrections Count": len(audit),
                    "Notes": d.get("notes", ""),
                    "Queries Succeeded": d.get("queries_succeeded", 0),
                    "Search Coverage": d.get("search_coverage_summary", ""),
                    "Executed Queries": q_str,
                    "Data Source": d.get("data_source", "Live YouTube Data API v3"),
                    "Empirical Finding Notice": "Observed research finding in search sample during specified period; not proof of complete platform-wide absence.",
                    "Created At": d.get("created_at", ""),
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_csv(df, sanitize_formulas=True, include_bom=True)

    @staticmethod
    def export_products_excel(products: List[Any], session_info: Optional[Dict[str, Any]] = None) -> bytes:
        """Export discovered products to styled XLSX bytes with formula sanitization."""
        rows = []
        for p in products:
            d = p.to_dict() if hasattr(p, "to_dict") else dict(p)
            bd = d.get("opportunity_score_breakdown", {}) or {}
            q_list = d.get("queries_executed", [])
            q_str = "; ".join(q_list) if isinstance(q_list, list) else str(q_list)
            audit = d.get("audit_trail", []) or []
            rows.append(
                {
                    "Product Name": d.get("product_name", ""),
                    "Category": d.get("category", ""),
                    "Brand": d.get("brand", ""),
                    "Model": d.get("model", ""),
                    "Priority": d.get("priority", "Medium"),
                    "Niche": d.get("niche", ""),
                    "Classification": d.get("classification", ""),
                    "Classification Reason": d.get("classification_reason", ""),
                    "Matching Videos Count": d.get("matching_videos_count", 0),
                    "Unique Channels Count": d.get("unique_channels_count", 0),
                    "Latest Video Date": d.get("latest_video_date", "") or "None",
                    "Opportunity Score": d.get("opportunity_score", 50.0),
                    "Confidence Level": d.get("confidence_level", "Low"),
                    "Scarcity Points (0-40)": bd.get("opportunity_factor", 0.0),
                    "Coverage Points (0-25)": bd.get("search_coverage_factor", 0.0),
                    "Identity Points (0-20)": bd.get("identity_confidence_factor", 0.0),
                    "Relevance Points (0-15)": bd.get("niche_relevance_factor", 0.0),
                    "Source": d.get("source", "Extracted Candidate"),
                    "Evidence URL": d.get("evidence_url", ""),
                    "Evidence Reference": d.get("evidence_reference", ""),
                    "Verification Status": d.get("verification_status", "Extracted Candidate"),
                    "Original Name": d.get("original_name", ""),
                    "User Corrections Count": len(audit),
                    "Notes": d.get("notes", ""),
                    "Queries Succeeded": d.get("queries_succeeded", 0),
                    "Search Coverage": d.get("search_coverage_summary", ""),
                    "Executed Queries": q_str,
                    "Data Source": d.get("data_source", "Live YouTube Data API v3"),
                    "Empirical Finding Notice": "Observed research finding in search sample during specified period; not proof of complete platform-wide absence.",
                    "Created At": d.get("created_at", ""),
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_excel(df, sanitize_formulas=True, sheet_name="Product Opportunities")

    @staticmethod
    def export_shortlist_csv(products: List[Any]) -> bytes:
        """Export shortlisted product opportunities to CSV bytes with UTF-8 BOM and formula sanitization."""
        rows = []
        for p in products:
            d = p.to_dict() if hasattr(p, "to_dict") else dict(p)
            aff = d.get("affiliate_data", {}) or {}
            bd = d.get("score_breakdown", {}) or {}
            chk = d.get("verification_checklist", {}) or {}
            sk = d.get("supporting_keywords", []) or []

            price_val = aff.get("price")
            curr_val = aff.get("currency", "USD")
            price_display = f"${price_val:,.2f} {curr_val}" if price_val is not None and price_val > 0 else "Unavailable"

            rows.append(
                {
                    "Product Name": d.get("product_name", ""),
                    "Brand": d.get("brand", "") or "Unavailable",
                    "Model": d.get("model", "") or "Unavailable",
                    "Niche": d.get("niche", ""),
                    "Category": d.get("category", "") or "Unavailable",
                    "Priority": d.get("priority", "Medium"),
                    "Workflow Status": d.get("status", "Shortlisted"),
                    "Intended Format": d.get("intended_format", "Product Review"),
                    "Opportunity Score (0-100)": d.get("opportunity_score", 50.0),
                    "Confidence Level": d.get("confidence_level", "Moderate Evidence"),
                    "Formula Version": d.get("formula_version", "v14.0"),
                    "Matching Videos Observed": d.get("matching_videos_count", 0),
                    "Queries Executed Count": d.get("queries_executed_count", 0),
                    "Identity Verified": "Yes" if chk.get("identity_verified") else "No",
                    "Specs Checked": "Yes" if chk.get("specs_checked") else "No",
                    "Merchant": aff.get("merchant", "") or "Needs verification",
                    "Price": price_display,
                    "Affiliate Program": aff.get("affiliate_program", "") or "Needs verification",
                    "Commission Info": aff.get("commission_info", "") or "Needs verification",
                    "Product Availability": aff.get("availability", "") or "Unknown",
                    "Affiliate Link": aff.get("affiliate_link", ""),
                    "Supporting Keywords": "; ".join(sk) if isinstance(sk, list) else str(sk),
                    "Notes": d.get("notes", ""),
                    "Notice": "Observed gap in search results; does not guarantee search volume or sales.",
                    "Created At": d.get("created_at", ""),
                    "Updated At": d.get("updated_at", ""),
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_csv(df, sanitize_formulas=True, include_bom=True)

    @staticmethod
    def export_shortlist_excel(products: List[Any]) -> bytes:
        """Export shortlisted product opportunities to styled XLSX bytes with formula sanitization."""
        rows = []
        for p in products:
            d = p.to_dict() if hasattr(p, "to_dict") else dict(p)
            aff = d.get("affiliate_data", {}) or {}
            chk = d.get("verification_checklist", {}) or {}
            sk = d.get("supporting_keywords", []) or []

            price_val = aff.get("price")
            curr_val = aff.get("currency", "USD")
            price_display = f"${price_val:,.2f} {curr_val}" if price_val is not None and price_val > 0 else "Unavailable"

            rows.append(
                {
                    "Product Name": d.get("product_name", ""),
                    "Brand": d.get("brand", "") or "Unavailable",
                    "Model": d.get("model", "") or "Unavailable",
                    "Niche": d.get("niche", ""),
                    "Category": d.get("category", "") or "Unavailable",
                    "Priority": d.get("priority", "Medium"),
                    "Workflow Status": d.get("status", "Shortlisted"),
                    "Intended Format": d.get("intended_format", "Product Review"),
                    "Opportunity Score (0-100)": d.get("opportunity_score", 50.0),
                    "Confidence Level": d.get("confidence_level", "Moderate Evidence"),
                    "Formula Version": d.get("formula_version", "v14.0"),
                    "Matching Videos Observed": d.get("matching_videos_count", 0),
                    "Queries Executed Count": d.get("queries_executed_count", 0),
                    "Identity Verified": "Yes" if chk.get("identity_verified") else "No",
                    "Specs Checked": "Yes" if chk.get("specs_checked") else "No",
                    "Merchant": aff.get("merchant", "") or "Needs verification",
                    "Price": price_display,
                    "Affiliate Program": aff.get("affiliate_program", "") or "Needs verification",
                    "Commission Info": aff.get("commission_info", "") or "Needs verification",
                    "Product Availability": aff.get("availability", "") or "Unknown",
                    "Affiliate Link": aff.get("affiliate_link", ""),
                    "Supporting Keywords": "; ".join(sk) if isinstance(sk, list) else str(sk),
                    "Notes": d.get("notes", ""),
                    "Notice": "Observed gap in search results; does not guarantee search volume or sales.",
                    "Created At": d.get("created_at", ""),
                    "Updated At": d.get("updated_at", ""),
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_excel(df, sanitize_formulas=True, sheet_name="Shortlisted Products")

    @staticmethod
    def export_affiliate_brief_csv(brief: Any) -> bytes:
        """Export an affiliate content brief to CSV bytes with formula sanitization."""
        d = brief.to_dict() if hasattr(brief, "to_dict") else dict(brief)
        p_cons = d.get("pros_and_cons", {}) or {}
        pros = "; ".join(p_cons.get("pros", []))
        cons = "; ".join(p_cons.get("cons", []))
        struct = d.get("video_structure", []) or []
        struct_str = " | ".join([f"{s.get('section', '')}: {s.get('notes', '')}" for s in struct])
        q_list = d.get("questions_to_investigate", []) or []
        feat_list = d.get("features_to_verify", []) or []
        titles = d.get("title_ideas", []) or []
        chk = d.get("fact_checking_checklist", []) or []
        chk_str = " | ".join([f"[{c.get('status', 'Pending')}] {c.get('task', '')}" for c in chk])

        rows = [
            {"Field": "Product Name", "Content": d.get("product_name", "")},
            {"Field": "Brand", "Content": d.get("brand", "") or "Unavailable"},
            {"Field": "Model", "Content": d.get("model", "") or "Unavailable"},
            {"Field": "Category", "Content": d.get("category", "") or "Unavailable"},
            {"Field": "Primary Keyword", "Content": d.get("primary_keyword", "")},
            {"Field": "Supporting Keywords", "Content": "; ".join(d.get("supporting_keywords", []))},
            {"Field": "Target Audience", "Content": d.get("target_audience", "")},
            {"Field": "Search Intent", "Content": d.get("search_intent", "")},
            {"Field": "Content Format", "Content": d.get("content_format", "")},
            {"Field": "Title Ideas", "Content": " | ".join(titles)},
            {"Field": "Video Outline & Structure", "Content": struct_str},
            {"Field": "Questions to Investigate", "Content": "; ".join(q_list)},
            {"Field": "Features to Verify", "Content": "; ".join(feat_list)},
            {"Field": "Objective Pros", "Content": pros},
            {"Field": "Objective Cons", "Content": cons},
            {"Field": "Call to Action", "Content": d.get("call_to_action", "")},
            {"Field": "FTC Affiliate Disclosure Reminder", "Content": d.get("affiliate_disclosure", "")},
            {"Field": "Fact-Checking Checklist", "Content": chk_str},
            {"Field": "Notes & Research Observations", "Content": d.get("brief_notes", "")},
            {"Field": "Ethical Guideline", "Content": "Do not claim personal testing or hands-on experience without verified direct use."},
        ]
        df = pd.DataFrame(rows)
        return ExportService.export_csv(df, sanitize_formulas=True, include_bom=True)

    @staticmethod
    def export_affiliate_brief_excel(brief: Any) -> bytes:
        """Export an affiliate content brief to styled XLSX bytes with formula sanitization."""
        d = brief.to_dict() if hasattr(brief, "to_dict") else dict(brief)
        p_cons = d.get("pros_and_cons", {}) or {}
        pros = "; ".join(p_cons.get("pros", []))
        cons = "; ".join(p_cons.get("cons", []))
        struct = d.get("video_structure", []) or []
        struct_str = " | ".join([f"{s.get('section', '')}: {s.get('notes', '')}" for s in struct])
        q_list = d.get("questions_to_investigate", []) or []
        feat_list = d.get("features_to_verify", []) or []
        titles = d.get("title_ideas", []) or []
        chk = d.get("fact_checking_checklist", []) or []
        chk_str = " | ".join([f"[{c.get('status', 'Pending')}] {c.get('task', '')}" for c in chk])

        rows = [
            {"Section": "Metadata", "Item": "Product Name", "Details": d.get("product_name", "")},
            {"Section": "Metadata", "Item": "Brand & Model", "Details": f"{d.get('brand', '')} {d.get('model', '')}".strip() or "Unavailable"},
            {"Section": "SEO", "Item": "Primary Keyword", "Details": d.get("primary_keyword", "")},
            {"Section": "SEO", "Item": "Supporting Keywords", "Details": "; ".join(d.get("supporting_keywords", []))},
            {"Section": "Strategy", "Item": "Target Audience", "Details": d.get("target_audience", "")},
            {"Section": "Strategy", "Item": "Search Intent", "Details": d.get("search_intent", "")},
            {"Section": "Strategy", "Item": "Content Format", "Details": d.get("content_format", "")},
            {"Section": "Production", "Item": "Title Ideas", "Details": " | ".join(titles)},
            {"Section": "Production", "Item": "Video Outline", "Details": struct_str},
            {"Section": "Fact Check", "Item": "Questions to Investigate", "Details": "; ".join(q_list)},
            {"Section": "Fact Check", "Item": "Features to Verify", "Details": "; ".join(feat_list)},
            {"Section": "Evaluation", "Item": "Objective Pros", "Details": pros},
            {"Section": "Evaluation", "Item": "Objective Cons", "Details": cons},
            {"Section": "Conversion", "Item": "Call to Action", "Details": d.get("call_to_action", "")},
            {"Section": "Compliance", "Item": "FTC Disclosure", "Details": d.get("affiliate_disclosure", "")},
            {"Section": "Compliance", "Item": "Fact-Checking Tasks", "Details": chk_str},
            {"Section": "Notes", "Item": "Brief Notes", "Details": d.get("brief_notes", "")},
            {"Section": "Ethics", "Item": "Ethical Reminder", "Details": "No unsupported claims of personal testing; disclose all affiliate relationships transparently."},
        ]
        df = pd.DataFrame(rows)
        return ExportService.export_excel(df, sanitize_formulas=True, sheet_name="Affiliate Content Brief")

    @staticmethod
    def export_comparison_matrix_csv(matrix_rows: List[Dict[str, Any]]) -> bytes:
        """Export product comparison matrix to CSV bytes with formula sanitization."""
        df = pd.DataFrame(matrix_rows)
        return ExportService.export_csv(df, sanitize_formulas=True, include_bom=True)

    @staticmethod
    def export_comparison_matrix_excel(matrix_rows: List[Dict[str, Any]]) -> bytes:
        """Export product comparison matrix to styled XLSX bytes with formula sanitization."""
        df = pd.DataFrame(matrix_rows)
        return ExportService.export_excel(df, sanitize_formulas=True, sheet_name="Opportunity Comparison")

    @staticmethod
    def export_monitor_runs_csv(runs: List[Any]) -> bytes:
        """Export research monitor run history to CSV bytes with formula sanitization and UTF-8 BOM."""
        rows = []
        for r in runs:
            d = r.to_dict() if hasattr(r, "to_dict") else dict(r)
            snap = d.get("config_snapshot", {}) or {}
            rows.append(
                {
                    "Run ID": d.get("id"),
                    "Run #": d.get("run_number", 1),
                    "Niche": d.get("niche", ""),
                    "Trigger Type": d.get("trigger_type", "Manual"),
                    "Status": d.get("status", "Completed"),
                    "Products Discovered": d.get("products_discovered", 0),
                    "Matching Videos Retrieved": d.get("matching_videos_retrieved", 0),
                    "Queries Attempted": d.get("queries_attempted", 0),
                    "Queries Succeeded": d.get("queries_succeeded", 0),
                    "Queries Failed": d.get("queries_failed", 0),
                    "Live API Requests": d.get("live_requests_count", 0),
                    "Cached Requests": d.get("cached_requests_count", 0),
                    "Provider": d.get("provider_name", "YouTube Data API v3"),
                    "Config Date Window (Days)": snap.get("date_window_days", 60),
                    "Config Request Budget": snap.get("request_budget", 15),
                    "Error Message": d.get("error_message") or "None",
                    "Started At": d.get("started_at", ""),
                    "Completed At": d.get("completed_at", "") or "In Progress",
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_csv(df, sanitize_formulas=True, include_bom=True)

    @staticmethod
    def export_monitor_runs_excel(runs: List[Any]) -> bytes:
        """Export research monitor run history to styled XLSX bytes with formula sanitization."""
        rows = []
        for r in runs:
            d = r.to_dict() if hasattr(r, "to_dict") else dict(r)
            snap = d.get("config_snapshot", {}) or {}
            rows.append(
                {
                    "Run ID": d.get("id"),
                    "Run #": d.get("run_number", 1),
                    "Niche": d.get("niche", ""),
                    "Trigger Type": d.get("trigger_type", "Manual"),
                    "Status": d.get("status", "Completed"),
                    "Products Discovered": d.get("products_discovered", 0),
                    "Matching Videos Retrieved": d.get("matching_videos_retrieved", 0),
                    "Queries Attempted": d.get("queries_attempted", 0),
                    "Queries Succeeded": d.get("queries_succeeded", 0),
                    "Queries Failed": d.get("queries_failed", 0),
                    "Live API Requests": d.get("live_requests_count", 0),
                    "Cached Requests": d.get("cached_requests_count", 0),
                    "Provider": d.get("provider_name", "YouTube Data API v3"),
                    "Config Date Window (Days)": snap.get("date_window_days", 60),
                    "Config Request Budget": snap.get("request_budget", 15),
                    "Error Message": d.get("error_message") or "None",
                    "Started At": d.get("started_at", ""),
                    "Completed At": d.get("completed_at", "") or "In Progress",
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_excel(df, sanitize_formulas=True, sheet_name="Research Runs")

    @staticmethod
    def export_snapshot_comparison_csv(comparison: Any) -> bytes:
        """Export snapshot comparison matrix and summary to CSV bytes with formula sanitization."""
        d = comparison.to_dict() if hasattr(comparison, "to_dict") else dict(comparison)
        matrix = d.get("comparison_matrix", []) or []
        rows = []
        for m in matrix:
            row_dict = dict(m)
            row_dict["Niche"] = d.get("niche", "")
            row_dict["Run A Date"] = d.get("run_a_date", "")[:10]
            row_dict["Run B Date"] = d.get("run_b_date", "")[:10]
            rows.append(row_dict)
        df = pd.DataFrame(rows)
        return ExportService.export_csv(df, sanitize_formulas=True, include_bom=True)

    @staticmethod
    def export_snapshot_comparison_excel(comparison: Any) -> bytes:
        """Export snapshot comparison matrix and summary to styled XLSX bytes with formula sanitization."""
        d = comparison.to_dict() if hasattr(comparison, "to_dict") else dict(comparison)
        matrix = d.get("comparison_matrix", []) or []
        rows = []
        for m in matrix:
            row_dict = dict(m)
            row_dict["Niche"] = d.get("niche", "")
            row_dict["Run A Date"] = d.get("run_a_date", "")[:10]
            row_dict["Run B Date"] = d.get("run_b_date", "")[:10]
            rows.append(row_dict)
        df = pd.DataFrame(rows)
        return ExportService.export_excel(df, sanitize_formulas=True, sheet_name="Snapshot Comparison")

    @staticmethod
    def export_monitor_alerts_csv(alerts: List[Any]) -> bytes:
        """Export change alerts with supporting evidence to CSV bytes with formula sanitization."""
        rows = []
        for a in alerts:
            d = a.to_dict() if hasattr(a, "to_dict") else dict(a)
            ev = d.get("evidence_data", {}) or {}
            rows.append(
                {
                    "Alert ID": d.get("id"),
                    "Severity": d.get("severity", "Info"),
                    "Alert Type": d.get("alert_type", ""),
                    "Niche": d.get("niche", ""),
                    "Product Name": d.get("product_name", "") or "Niche Wide",
                    "Title": d.get("title", ""),
                    "Message": d.get("message", ""),
                    "Evidence Details": json.dumps(ev),
                    "Dismissed": "Yes" if d.get("is_dismissed") else "No",
                    "Created At": d.get("created_at", ""),
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_csv(df, sanitize_formulas=True, include_bom=True)

    @staticmethod
    def export_monitor_alerts_excel(alerts: List[Any]) -> bytes:
        """Export change alerts with supporting evidence to styled XLSX bytes with formula sanitization."""
        rows = []
        for a in alerts:
            d = a.to_dict() if hasattr(a, "to_dict") else dict(a)
            ev = d.get("evidence_data", {}) or {}
            rows.append(
                {
                    "Alert ID": d.get("id"),
                    "Severity": d.get("severity", "Info"),
                    "Alert Type": d.get("alert_type", ""),
                    "Niche": d.get("niche", ""),
                    "Product Name": d.get("product_name", "") or "Niche Wide",
                    "Title": d.get("title", ""),
                    "Message": d.get("message", ""),
                    "Evidence Details": json.dumps(ev),
                    "Dismissed": "Yes" if d.get("is_dismissed") else "No",
                    "Created At": d.get("created_at", ""),
                }
            )
        df = pd.DataFrame(rows)
        return ExportService.export_excel(df, sanitize_formulas=True, sheet_name="Change Alerts")

    @staticmethod
    def export_historical_trends_csv(timeline_rows: List[Dict[str, Any]]) -> bytes:
        """Export historical timeline trends to CSV bytes with formula sanitization."""
        df = pd.DataFrame(timeline_rows)
        return ExportService.export_csv(df, sanitize_formulas=True, include_bom=True)

    @staticmethod
    def export_historical_trends_excel(timeline_rows: List[Dict[str, Any]]) -> bytes:
        """Export historical timeline trends to styled XLSX bytes with formula sanitization."""
        df = pd.DataFrame(timeline_rows)
        return ExportService.export_excel(df, sanitize_formulas=True, sheet_name="Historical Trends")


# Phase 15 Module-level export aliases
export_monitor_runs_csv = ExportService.export_monitor_runs_csv
export_monitor_runs_excel = ExportService.export_monitor_runs_excel
export_snapshot_comparison_csv = ExportService.export_snapshot_comparison_csv
export_snapshot_comparison_excel = ExportService.export_snapshot_comparison_excel
export_monitor_alerts_csv = ExportService.export_monitor_alerts_csv
export_monitor_alerts_excel = ExportService.export_monitor_alerts_excel
export_historical_trends_csv = ExportService.export_historical_trends_csv
export_historical_trends_excel = ExportService.export_historical_trends_excel

__all__ = [
    "ExportService",
    "sanitize_formula_injection",
    "sanitize_dataframe_for_export",
    "generate_export_filename",
    "export_monitor_runs_csv",
    "export_monitor_runs_excel",
    "export_snapshot_comparison_csv",
    "export_snapshot_comparison_excel",
    "export_monitor_alerts_csv",
    "export_monitor_alerts_excel",
    "export_historical_trends_csv",
    "export_historical_trends_excel",
]
