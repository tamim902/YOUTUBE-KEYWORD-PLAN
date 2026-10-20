"""Phase 10 automated test suite: YouTube Video SERP & Competitor Inspector.

Covers:
1. SERP Domain models (CompetitorVideoItem, KeywordSerpAnalysis, QuotaUsageLog).
2. UnconnectedSerpMetricsProvider honest offline status (zero synthetic metrics).
3. ISO 8601 duration parser (PT#H#M#S, PT#M#S, PT#S).
4. Relative age calculator and human-readable format.
5. Opportunity score formula and explainable criteria.
6. SQLite application settings persistence.
7. Local SERP caching, normalization, and TTL expiry.
8. Daily API quota tracking and 10,000 unit limit enforcement.
9. YouTubeDataApiSerpProvider API key verification.
10. Batched search and video details execution (101 quota units).
11. Cache hit optimization (0 quota units).
12. SERP CSV export with UTF-8 BOM and formula sanitization.
13. SERP Excel export with openpyxl styling and formula protection.
14. Multilingual & Bengali Unicode SERP support.
"""

import io
import json
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import MagicMock

import openpyxl
import pytest
import requests

from src.models.serp import (
    CompetitorVideoItem,
    KeywordSerpAnalysis,
    KeywordValidationSummary,
    QuotaUsageLog,
    TitlePatternAnalysis,
)
from src.providers.serp import (
    SerpMetrics,
    UnconnectedSerpMetricsProvider,
    YouTubeDataApiSerpProvider,
    analyze_title_patterns,
    calculate_age_distribution,
    calculate_opportunity_score,
    calculate_relative_age,
    calculate_view_count_distribution,
    default_serp_provider,
    extract_common_title_phrases,
    generate_keyword_validation_summary,
    parse_iso_duration,
)
from src.storage.database import DatabaseManager
from src.storage.export import ExportService


# ------------------------------------------------------------------------------
# 1. SERP Models
# ------------------------------------------------------------------------------

def test_serp_models_instantiation_and_serialization() -> None:
    """Verify CompetitorVideoItem, KeywordSerpAnalysis, and QuotaUsageLog models."""
    item = CompetitorVideoItem(
        video_id="abc123xyz",
        title="Best Car Wax 2026",
        channel_id="ch_1",
        channel_title="Auto Lab",
        publish_date="2026-01-10T12:00:00Z",
        view_count=50000,
        like_count=2500,
        comment_count=180,
        duration="14:22",
        tags=["car wax", "detailing"],
        rank=1,
    )
    assert item.video_url == "https://www.youtube.com/watch?v=abc123xyz"
    d = item.to_dict()
    assert d["video_id"] == "abc123xyz"
    assert d["view_count"] == 50000

    analysis = KeywordSerpAnalysis(
        keyword="best car wax",
        videos=[item],
        total_videos_analyzed=1,
        average_views=50000.0,
        median_views=50000.0,
        opportunity_score=75.0,
        opportunity_label="High Opportunity",
        top_tags=[("car wax", 1), ("detailing", 1)],
    )
    ad = analysis.to_dict()
    assert ad["keyword"] == "best car wax"
    assert len(ad["videos"]) == 1
    assert ad["videos"][0]["video_id"] == "abc123xyz"

    quota = QuotaUsageLog(date="2026-10-18", units_used=202, daily_limit=10000)
    assert quota.units_remaining == 9798
    assert round(quota.percentage_used, 2) == 2.02


# ------------------------------------------------------------------------------
# 2. Honest Unconnected Provider
# ------------------------------------------------------------------------------

def test_unconnected_provider_honest_status() -> None:
    """Verify unconnected provider reports honest unavailable status without synthetic metrics."""
    provider = UnconnectedSerpMetricsProvider()
    assert not provider.is_connected()

    metrics = provider.get_metrics("python tutorial")
    assert not metrics.is_available
    assert "Not available — no verified data source connected" in metrics.status_message
    assert metrics.search_volume is None
    assert metrics.competition_score is None
    assert metrics.average_views is None


# ------------------------------------------------------------------------------
# 3. ISO 8601 Duration Parsing
# ------------------------------------------------------------------------------

@pytest.mark.parametrize(
    "iso_str, expected_fmt, expected_sec",
    [
        ("PT14M22S", "14:22", 862),
        ("PT1H2M10S", "1:02:10", 3730),
        ("PT45S", "0:45", 45),
        ("PT1H", "1:00:00", 3600),
        ("PT2M", "2:00", 120),
        ("P1DT2H30M", "26:30:00", 95400),
        ("", "0:00", 0),
        ("INVALID", "0:00", 0),
    ],
)
def test_parse_iso_duration(iso_str: str, expected_fmt: str, expected_sec: int) -> None:
    """Verify ISO 8601 duration parser accurately converts YouTube formats to human strings and seconds."""
    fmt, sec = parse_iso_duration(iso_str)
    assert fmt == expected_fmt
    assert sec == expected_sec


# ------------------------------------------------------------------------------
# 4. Relative Age Calculation
# ------------------------------------------------------------------------------

def test_calculate_relative_age() -> None:
    """Verify relative age produces human-readable outputs and elapsed day counts."""
    now_iso = datetime.now(timezone.utc).isoformat()
    txt, days = calculate_relative_age(now_iso)
    assert txt == "Today"
    assert days == 0

    txt_inv, days_inv = calculate_relative_age("invalid-date")
    assert txt_inv == "Unknown"
    assert days_inv == 0


# ------------------------------------------------------------------------------
# 5. Opportunity Score Calculation
# ------------------------------------------------------------------------------

def test_calculate_opportunity_score_empty() -> None:
    """Empty video list defaults to moderate 50.0 score."""
    score, label = calculate_opportunity_score([])
    assert score == 50.0
    assert "Moderate" in label


def test_calculate_opportunity_score_high_opportunity() -> None:
    """High views, older ranking videos, and accessible competitors yield high opportunity."""
    videos = [
        CompetitorVideoItem(
            video_id=f"v_{i}",
            title=f"Video {i}",
            channel_id=f"ch_{i}",
            channel_title=f"Channel {i}",
            publish_date="2023-01-01T00:00:00Z",  # Over 2 years old
            view_count=25000 if i % 2 == 0 else 120000,  # Half under 50k views
        )
        for i in range(10)
    ]
    score, label = calculate_opportunity_score(videos)
    assert score >= 75.0
    assert "High Opportunity" in label


def test_calculate_opportunity_score_saturated() -> None:
    """Videos dominated by mega view counts with zero accessible small competitors."""
    videos = [
        CompetitorVideoItem(
            video_id=f"v_{i}",
            title=f"Mega Video {i}",
            channel_id=f"ch_{i}",
            channel_title=f"Mega Channel {i}",
            publish_date="2026-09-01T00:00:00Z",  # Very fresh
            view_count=5000000,  # 5 million views each
        )
        for i in range(10)
    ]
    score, label = calculate_opportunity_score(videos)
    assert score <= 60.0


# ------------------------------------------------------------------------------
# 6. SQLite Settings Persistence
# ------------------------------------------------------------------------------

def test_settings_persistence(tmp_path: Path) -> None:
    """Verify app_settings table saves, retrieves, and updates configuration."""
    db_file = tmp_path / "test_settings.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    assert mgr.get_setting("youtube_api_key") is None
    assert mgr.get_setting("youtube_api_key", "default_val") == "default_val"

    mgr.set_setting("youtube_api_key", "AIzaSyTestKey123")
    assert mgr.get_setting("youtube_api_key") == "AIzaSyTestKey123"

    mgr.set_setting("youtube_api_key", "AIzaSyUpdatedKey456")
    assert mgr.get_setting("youtube_api_key") == "AIzaSyUpdatedKey456"


# ------------------------------------------------------------------------------
# 7. Local SERP Caching & Normalization
# ------------------------------------------------------------------------------

def test_serp_caching_and_normalization(tmp_path: Path) -> None:
    """Verify SERP caching, case-insensitive keyword normalization, and TTL pruning."""
    db_file = tmp_path / "test_serp_cache.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    assert mgr.get_cached_serp("best car wax", "en", "US") is None

    payload = {
        "keyword": "best car wax",
        "videos": [{"video_id": "v1", "title": "Wax Guide", "channel_id": "c1", "channel_title": "CarLab", "publish_date": "2026-01-01T00:00:00Z", "view_count": 1000}],
        "average_views": 1000.0,
    }
    mgr.set_cached_serp("  BEST   car   WAX  ", "en", "US", payload)

    cached = mgr.get_cached_serp("best car wax", "en", "US")
    assert cached is not None
    assert cached["keyword"] == "best car wax"
    assert len(cached["videos"]) == 1

    # TTL expiry check
    expired = mgr.get_cached_serp("best car wax", "en", "US", max_age_seconds=0)
    assert expired is None

    cleared = mgr.clear_serp_cache()
    assert cleared == 1
    assert mgr.get_cached_serp("best car wax", "en", "US") is None


# ------------------------------------------------------------------------------
# 8. Daily Quota Tracking
# ------------------------------------------------------------------------------

def test_api_quota_tracking(tmp_path: Path) -> None:
    """Verify API quota usage tracking and daily counter reset."""
    db_file = tmp_path / "test_quota.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    today_str = "2026-10-18"
    assert mgr.get_daily_quota_used(today_str) == 0

    total = mgr.record_quota_usage(101, today_str)
    assert total == 101

    total2 = mgr.record_quota_usage(101, today_str)
    assert total2 == 202
    assert mgr.get_daily_quota_used(today_str) == 202

    mgr.reset_daily_quota(today_str)
    assert mgr.get_daily_quota_used(today_str) == 0


# ------------------------------------------------------------------------------
# 9. YouTube SERP Provider Key Verification
# ------------------------------------------------------------------------------

def test_youtube_serp_provider_verify_api_key() -> None:
    """Verify YouTubeDataApiSerpProvider.verify_api_key with mocked HTTP responses."""
    mock_session = MagicMock(spec=requests.Session)

    # 200 OK
    resp_200 = MagicMock()
    resp_200.status_code = 200
    mock_session.get.return_value = resp_200

    provider = YouTubeDataApiSerpProvider(session=mock_session)
    ok, msg = provider.verify_api_key("valid_key_123")
    assert ok
    assert "successfully verified" in msg

    # 400 Bad Request
    resp_400 = MagicMock()
    resp_400.status_code = 400
    resp_400.json.return_value = {"error": {"message": "API key not valid."}}
    mock_session.get.return_value = resp_400

    bad_ok, bad_msg = provider.verify_api_key("invalid_key_456")
    assert not bad_ok
    assert "API key not valid" in bad_msg

    # Empty key
    empty_ok, empty_msg = provider.verify_api_key("")
    assert not empty_ok


# ------------------------------------------------------------------------------
# 10. Live SERP Retrieval & Batching
# ------------------------------------------------------------------------------

def test_youtube_serp_provider_analyze_keyword_live(tmp_path: Path) -> None:
    """Verify 2-step batch query (search.list + batched videos.list) costing 101 units."""
    db_file = tmp_path / "test_serp_live.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    mock_session = MagicMock(spec=requests.Session)

    # Search response
    search_resp = MagicMock()
    search_resp.status_code = 200
    search_resp.json.return_value = {
        "items": [
            {"id": {"videoId": "vid_1"}},
            {"id": {"videoId": "vid_2"}},
        ]
    }

    # Batched videos response
    videos_resp = MagicMock()
    videos_resp.status_code = 200
    videos_resp.json.return_value = {
        "items": [
            {
                "id": "vid_1",
                "snippet": {
                    "title": "Top Waxing Tips",
                    "channelId": "ch_1",
                    "channelTitle": "Auto Detailers",
                    "publishedAt": "2025-05-01T10:00:00Z",
                    "tags": ["car wax", "detailing", "shine"],
                    "thumbnails": {"medium": {"url": "https://img.youtube.com/vi/vid_1/mqdefault.jpg"}},
                },
                "statistics": {
                    "viewCount": "85000",
                    "likeCount": "4200",
                    "commentCount": "310",
                },
                "contentDetails": {"duration": "PT15M30S"},
            },
            {
                "id": "vid_2",
                "snippet": {
                    "title": "Ceramic vs Wax",
                    "channelId": "ch_2",
                    "channelTitle": "Detail Pro",
                    "publishedAt": "2024-01-15T08:00:00Z",
                    "tags": ["car wax", "ceramic coating"],
                    "thumbnails": {"medium": {"url": "https://img.youtube.com/vi/vid_2/mqdefault.jpg"}},
                },
                "statistics": {
                    "viewCount": "190000",
                    "likeCount": "9100",
                    "commentCount": "720",
                },
                "contentDetails": {"duration": "PT8M45S"},
            },
        ]
    }

    mock_session.get.side_effect = [search_resp, videos_resp]

    provider = YouTubeDataApiSerpProvider(
        api_key="valid_api_key_xyz",
        db_manager=mgr,
        session=mock_session,
    )

    analysis = provider.analyze_keyword("best car wax", use_cache=False)
    assert analysis.total_videos_analyzed == 2
    assert analysis.videos[0].video_id == "vid_1"
    assert analysis.videos[0].duration == "15:30"
    assert analysis.videos[0].view_count == 85000
    assert analysis.average_views == 137500.0
    assert analysis.median_views == 137500.0
    assert len(analysis.top_tags) > 0
    assert analysis.top_tags[0][0] == "car wax"  # Most common tag
    assert analysis.top_tags[0][1] == 2

    # Verify quota recorded: exactly 101 units
    assert mgr.get_daily_quota_used() == 101


# ------------------------------------------------------------------------------
# 11. Cache Hit Consumes 0 Quota Units
# ------------------------------------------------------------------------------

def test_serp_cache_hit_zero_quota(tmp_path: Path) -> None:
    """Verify subsequent lookups read from local cache and consume 0 quota units."""
    db_file = tmp_path / "test_serp_cached_zero.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    provider = YouTubeDataApiSerpProvider(api_key="valid_key", db_manager=mgr)

    # Pre-populate cache
    cached_payload = {
        "keyword": "python tutorial",
        "videos": [
            {
                "video_id": "py1",
                "title": "Python for Beginners",
                "channel_id": "ch1",
                "channel_title": "Code Academy",
                "publish_date": "2025-01-01T00:00:00Z",
                "view_count": 500000,
                "like_count": 30000,
                "comment_count": 1200,
                "duration": "1:00:00",
                "tags": ["python", "coding"],
            }
        ],
        "total_videos_analyzed": 1,
        "average_views": 500000.0,
        "median_views": 500000.0,
        "opportunity_score": 68.0,
        "opportunity_label": "Moderate Opportunity",
        "top_tags": [["python", 1]],
    }
    mgr.set_cached_serp("python tutorial", "en", "US", cached_payload)

    # Query with use_cache=True
    analysis = provider.analyze_keyword("python tutorial", "en", "US", use_cache=True)
    assert analysis.is_cached
    assert "0 quota units used" in analysis.status_message
    assert mgr.get_daily_quota_used() == 0  # 0 units consumed!


# ------------------------------------------------------------------------------
# 12. Quota Limit Enforcement
# ------------------------------------------------------------------------------

def test_serp_quota_limit_enforcement(tmp_path: Path) -> None:
    """Verify provider halts requests once daily 10,000 units are exhausted."""
    db_file = tmp_path / "test_serp_cap.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    # Artificially consume 10,000 units
    mgr.record_quota_usage(10000)

    provider = YouTubeDataApiSerpProvider(api_key="valid_key", db_manager=mgr)
    analysis = provider.analyze_keyword("best wax", use_cache=False)

    assert analysis.total_videos_analyzed == 0
    assert "Daily API quota limit (10,000 units) reached" in analysis.status_message


# ------------------------------------------------------------------------------
# 13. CSV and Excel SERP Exports
# ------------------------------------------------------------------------------

def test_serp_export_csv_and_excel() -> None:
    """Verify CSV (UTF-8 BOM, formula-sanitized) and Excel exports for SERP analysis."""
    item = CompetitorVideoItem(
        video_id="v123",
        title="=SUM(A1)+Best Car Wax @Home",  # Formula injection attempt
        channel_id="c1",
        channel_title="Auto Pro",
        publish_date="2026-01-01T00:00:00Z",
        view_count=12000,
        like_count=500,
        comment_count=40,
        duration="10:00",
        tags=["car wax", "review"],
        rank=1,
    )
    analysis = KeywordSerpAnalysis(
        keyword="best car wax",
        videos=[item],
        total_videos_analyzed=1,
        average_views=12000.0,
        median_views=12000.0,
        opportunity_score=82.0,
        opportunity_label="High Opportunity",
        top_tags=[("car wax", 1)],
    )

    # 1. CSV Export
    csv_bytes = ExportService.export_serp_csv(analysis)
    assert csv_bytes.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM
    csv_str = csv_bytes.decode("utf-8-sig")
    assert "'=SUM(A1)" in csv_str  # Formula sanitized
    assert "High Opportunity" in csv_str

    # 2. Excel Export
    excel_bytes = ExportService.export_serp_excel(analysis)
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "SERP Competitors" in wb.sheetnames
    ws = wb["SERP Competitors"]
    # Check that formula is neutralized
    cell_val = None
    for row in ws.iter_rows(values_only=True):
        for val in row:
            if val and "SUM(A1)" in str(val):
                cell_val = str(val)
    assert cell_val is not None
    assert cell_val.startswith("'=")


# ------------------------------------------------------------------------------
# 14. Multilingual & Bengali Unicode Support
# ------------------------------------------------------------------------------

def test_bangla_unicode_serp_support(tmp_path: Path) -> None:
    """Verify Bengali Unicode keywords are preserved throughout SERP caching and export."""
    db_file = tmp_path / "test_bangla_serp.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    bn_keyword = "গরুর মাংস রান্নার রেসিপি"
    item = CompetitorVideoItem(
        video_id="bn_v1",
        title="ঐতিহ্যবাহী স্বাদে গরুর মাংসের ভুনা রেসিপি",
        channel_id="ch_bn",
        channel_title="গ্রামীণ রান্নাঘর",
        publish_date="2025-06-15T12:00:00Z",
        view_count=450000,
        tags=["গরুর মাংস", "রান্না", "রেসিপি"],
        rank=1,
    )
    analysis = KeywordSerpAnalysis(
        keyword=bn_keyword,
        language="bn",
        country="BD",
        videos=[item],
        total_videos_analyzed=1,
        average_views=450000.0,
        median_views=450000.0,
        opportunity_score=80.0,
        opportunity_label="High Opportunity",
        top_tags=[("গরুর মাংস", 1), ("রান্না", 1)],
    )

    # Test cache storage & retrieval
    mgr.set_cached_serp(bn_keyword, "bn", "BD", analysis.to_dict())
    cached = mgr.get_cached_serp(bn_keyword, "bn", "BD")
    assert cached is not None
    assert cached["keyword"] == bn_keyword
    assert cached["videos"][0]["title"] == "ঐতিহ্যবাহী স্বাদে গরুর মাংসের ভুনা রেসিপি"

    # Test CSV export has correct Bangla characters
    csv_bytes = ExportService.export_serp_csv(analysis)
    csv_str = csv_bytes.decode("utf-8-sig")
    assert "গরুর মাংস রান্নার রেসিপি" in csv_str
    assert "গ্রামীণ রান্নাঘর" in csv_str


# ------------------------------------------------------------------------------
# 15. Keyword Validation Summary Model
# ------------------------------------------------------------------------------

def test_keyword_validation_summary_model() -> None:
    """Verify KeywordValidationSummary model instantiation and dictionary serialization."""
    summary = KeywordValidationSummary(
        keyword="best car wax",
        total_videos=10,
        metadata_coverage={"view_count_pct": 100.0, "like_count_pct": 90.0},
        unique_channels_count=8,
        view_count_distribution={"min": 5000, "median": 45000.0, "max": 250000},
        age_distribution={"recent_under_30d": 2, "legacy_over_365d": 5},
        common_title_phrases=[("car wax", 5), ("tested reviewed", 3)],
        potential_content_angles=[{"angle": "Fresh Guide", "rationale": "5 videos > 1 year old"}],
        missing_data_limitations=["Sample limited to 10 top rankers"],
    )
    d = summary.to_dict()
    assert d["keyword"] == "best car wax"
    assert d["total_videos"] == 10
    assert d["unique_channels_count"] == 8
    assert d["metadata_coverage"]["view_count_pct"] == 100.0
    assert len(d["potential_content_angles"]) == 1


# ------------------------------------------------------------------------------
# 16. Title Pattern Analysis Model
# ------------------------------------------------------------------------------

def test_title_pattern_analysis_model() -> None:
    """Verify TitlePatternAnalysis model instantiation and dictionary serialization."""
    analysis = TitlePatternAnalysis(
        format_distribution={"Best & Top Lists": 4, "Reviews & Tests": 3},
        sample_size=10,
        detected_patterns=[{"format": "Best & Top Lists", "count": 4, "percentage": 40.0}],
        suggested_title_ideas=["Top 5 Best Car Wax Options Tested (2026)"],
        target_audience="Beginners",
        preferred_format="Reviews & Tests",
    )
    d = analysis.to_dict()
    assert d["sample_size"] == 10
    assert d["format_distribution"]["Best & Top Lists"] == 4
    assert d["target_audience"] == "Beginners"
    assert len(d["suggested_title_ideas"]) == 1


# ------------------------------------------------------------------------------
# 17. Distribution Calculations & Phrase Extraction
# ------------------------------------------------------------------------------

def test_view_count_and_age_distributions() -> None:
    """Verify view count and publication date distribution statistics."""
    # 1. Empty views
    empty_dist = calculate_view_count_distribution([])
    assert empty_dist["sample_size"] == 0
    assert empty_dist["median"] == 0.0

    # 2. Populated views
    views = [1000, 5000, 15000, 50000, 200000]
    dist = calculate_view_count_distribution(views)
    assert dist["sample_size"] == 5
    assert dist["min"] == 1000
    assert dist["max"] == 200000
    assert dist["median"] == 15000.0
    assert dist["average"] == 54200.0

    # 3. Age brackets
    ages = [5, 25, 45, 90, 200, 400, 800]
    age_dist = calculate_age_distribution(ages)
    assert age_dist["recent_under_30d"] == 2  # 5, 25
    assert age_dist["moderate_30_to_180d"] == 2  # 45, 90
    assert age_dist["established_180_to_365d"] == 1  # 200
    assert age_dist["legacy_over_365d"] == 2  # 400, 800


def test_extract_common_title_phrases() -> None:
    """Verify common 2-word and 3-word n-gram extraction from competitor titles."""
    titles = [
        "Best Car Wax for Black Cars Tested",
        "Top Rated Car Wax for Black Cars Review",
        "Best Car Wax Comparison for Black Cars",
    ]
    phrases = extract_common_title_phrases(titles, seed_keyword="best car wax")
    phrase_names = [p[0] for p in phrases]
    # "black cars" occurs in all 3 titles
    assert any("black cars" in p for p in phrase_names)


def test_analyze_title_patterns_and_suggestions() -> None:
    """Verify title pattern recognition and original suggestion generation."""
    titles = [
        "Top 10 Best Car Wax 2026",
        "Meguiar's Review and Test",
        "Ceramic Wax vs Carnauba Wax",
        "How to Apply Car Wax Like a Pro",
        "How to Fix Paint Swirls with Wax",
        "Car Wax for Beginners",
    ]
    pattern_result = analyze_title_patterns(
        titles=titles,
        keyword="car wax",
        target_audience="Beginners",
        preferred_format="How-To & Guides",
    )
    assert pattern_result.sample_size == 6
    assert pattern_result.format_distribution["Best & Top Lists"] >= 1
    assert pattern_result.format_distribution["Reviews & Tests"] >= 1
    assert pattern_result.format_distribution["Comparisons (VS)"] >= 1
    assert pattern_result.format_distribution["How-To & Guides"] >= 1
    assert pattern_result.format_distribution["Problem Solving"] >= 1
    assert pattern_result.format_distribution["Audience Specific"] >= 1
    assert len(pattern_result.suggested_title_ideas) >= 5
    assert any("Car Wax" in s for s in pattern_result.suggested_title_ideas)


def test_generate_keyword_validation_summary() -> None:
    """Verify 10-point evidence-based keyword validation summary generation."""
    # 1. Empty videos
    empty_val = generate_keyword_validation_summary("test query", [])
    assert empty_val.total_videos == 0
    assert len(empty_val.missing_data_limitations) > 0

    # 2. Populated videos
    videos = [
        CompetitorVideoItem(
            video_id=f"v_{i}",
            title=f"Best Car Wax {i}",
            channel_id=f"ch_{i % 3}",
            channel_title=f"Channel {i % 3}",
            publish_date="2024-01-01T00:00:00Z",
            view_count=(i + 1) * 10000,
            like_count=(i + 1) * 500,
            duration="10:00",
            tags=["car wax", "detailing"],
            description="Detailed review of car wax products.",
            rank=i + 1,
        )
        for i in range(5)
    ]
    val = generate_keyword_validation_summary("best car wax", videos)
    assert val.total_videos == 5
    assert val.unique_channels_count == 3  # 0, 1, 2
    assert val.metadata_coverage["view_count_pct"] == 100.0
    assert val.metadata_coverage["tags_pct"] == 100.0
    assert len(val.potential_content_angles) > 0
    assert len(val.missing_data_limitations) >= 3


# ------------------------------------------------------------------------------
# 18. SQLite Persistence for Video Research Sessions & Associations
# ------------------------------------------------------------------------------

def test_video_research_session_persistence(tmp_path: Path) -> None:
    """Verify saving, browsing, fetching, keyword association, and deleting research sessions."""
    db_file = tmp_path / "test_research_sessions.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    v1 = CompetitorVideoItem(
        video_id="vid_pers_1",
        title="Best Car Wax Guide",
        channel_id="ch_auto",
        channel_title="Auto Hub",
        publish_date="2026-02-01T00:00:00Z",
        view_count=85000,
        like_count=4200,
        comment_count=310,
        duration="12:30",
        tags=["car wax", "detailing", "auto"],
        rank=1,
    )
    v2 = CompetitorVideoItem(
        video_id="vid_pers_2",
        title="Ceramic Coating vs Wax",
        channel_id="ch_shine",
        channel_title="Shine Pro",
        publish_date="2026-01-15T00:00:00Z",
        view_count=120000,
        like_count=6100,
        comment_count=480,
        duration="18:45",
        tags=["ceramic", "wax"],
        rank=2,
    )

    analysis = KeywordSerpAnalysis(
        keyword="best car wax",
        language="en",
        country="US",
        videos=[v1, v2],
        total_videos_analyzed=2,
        average_views=102500.0,
        median_views=102500.0,
        opportunity_score=78.5,
        opportunity_label="High Opportunity",
        top_tags=[("car wax", 2), ("ceramic", 1)],
    )

    # 1. Save session
    session_id = mgr.save_video_research_session(
        analysis=analysis,
        provider="YouTube Data API v3",
        associated_keywords=["best car wax", "car wax review"],
    )
    assert session_id > 0

    # 2. List sessions
    sessions = mgr.get_video_research_sessions(limit=10)
    assert len(sessions) == 1
    assert sessions[0]["id"] == session_id
    assert sessions[0]["keyword"] == "best car wax"
    assert sessions[0]["total_videos"] == 2
    assert sessions[0]["metrics"]["median_views"] == 102500.0

    # 3. Get single session with items
    sess_detail = mgr.get_video_research_session(session_id)
    assert sess_detail is not None
    assert sess_detail["keyword"] == "best car wax"
    assert len(sess_detail["videos"]) == 2
    assert sess_detail["videos"][0]["video_id"] == "vid_pers_1"
    assert sess_detail["videos"][0]["tags"] == ["car wax", "detailing", "auto"]
    assert sess_detail["videos"][0]["channel_url"] == "https://www.youtube.com/channel/ch_auto"

    # 4. Check keyword associations
    kws_for_v1 = mgr.get_keywords_for_video("vid_pers_1")
    assert "best car wax" in kws_for_v1
    assert "car wax review" in kws_for_v1

    # 5. Deduplication across sessions: saving v1 again with another keyword
    analysis2 = KeywordSerpAnalysis(
        keyword="top car waxes",
        videos=[v1],
        total_videos_analyzed=1,
    )
    session_id2 = mgr.save_video_research_session(
        analysis=analysis2,
        associated_keywords=["top car waxes"],
    )
    assert session_id2 > session_id

    # Video 1 now associated with 3 keywords
    kws_updated = mgr.get_keywords_for_video("vid_pers_1")
    assert len(kws_updated) >= 2
    assert "top car waxes" in kws_updated

    # 6. Delete session 1
    deleted = mgr.delete_video_research_session(session_id)
    assert deleted is True
    assert mgr.get_video_research_session(session_id) is None
    # Session 2 still intact
    assert mgr.get_video_research_session(session_id2) is not None


# ------------------------------------------------------------------------------
# 19. Extended Provider Parameters: Order & Date Filtering
# ------------------------------------------------------------------------------

def test_youtube_data_api_serp_provider_extended_parameters(tmp_path: Path) -> None:
    """Verify analyze_keyword correctly passes order, publishedAfter, and generates validation summary."""
    db_file = tmp_path / "test_provider_extended.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    mock_session = MagicMock()

    # Search response mock
    mock_search_resp = MagicMock()
    mock_search_resp.status_code = 200
    mock_search_resp.json.return_value = {
        "items": [
            {"id": {"videoId": "test_vid_1"}},
            {"id": {"videoId": "test_vid_2"}},
        ]
    }

    # Videos list response mock
    mock_videos_resp = MagicMock()
    mock_videos_resp.status_code = 200
    mock_videos_resp.json.return_value = {
        "items": [
            {
                "id": "test_vid_1",
                "snippet": {
                    "title": "Top 10 Best Car Wax 2026",
                    "channelId": "ch_1",
                    "channelTitle": "Auto Testing",
                    "publishedAt": "2026-02-10T12:00:00Z",
                    "description": "Comprehensive review of 10 car waxes.",
                    "tags": ["car wax", "detailing"],
                    "thumbnails": {"medium": {"url": "https://img.youtube.com/vi/test_vid_1/mqdefault.jpg"}},
                },
                "statistics": {
                    "viewCount": "75000",
                    "likeCount": "3200",
                    "commentCount": "210",
                },
                "contentDetails": {
                    "duration": "PT15M30S",
                },
            },
            {
                "id": "test_vid_2",
                "snippet": {
                    "title": "How to Wax Your Car for Beginners",
                    "channelId": "ch_2",
                    "channelTitle": "DIY Garage",
                    "publishedAt": "2026-01-20T10:00:00Z",
                    "description": "Step by step beginner guide.",
                    "tags": ["car wax", "diy", "beginners"],
                    "thumbnails": {"medium": {"url": "https://img.youtube.com/vi/test_vid_2/mqdefault.jpg"}},
                },
                "statistics": {
                    "viewCount": "42000",
                    "likeCount": "1900",
                    "commentCount": "140",
                },
                "contentDetails": {
                    "duration": "PT10M15S",
                },
            },
        ]
    }

    mock_session.get.side_effect = [mock_search_resp, mock_videos_resp]

    provider = YouTubeDataApiSerpProvider(
        api_key="AIzaSyTestKey_12345",
        db_manager=mgr,
        session=mock_session,
    )

    published_after = "2026-01-01T00:00:00Z"
    analysis = provider.analyze_keyword(
        keyword="best car wax",
        language="en",
        country="US",
        max_results=10,
        order="viewCount",
        published_after=published_after,
        target_audience="Beginners",
        preferred_format="How-To & Guides",
        use_cache=False,
    )

    # Check search request arguments
    first_call_args, first_call_kwargs = mock_session.get.call_args_list[0]
    params = first_call_kwargs["params"]
    assert params["order"] == "viewCount"
    assert params["publishedAfter"] == published_after
    assert params["q"] == "best car wax"

    # Check returned analysis
    assert analysis.total_videos_analyzed == 2
    assert analysis.videos[0].channel_url == "https://www.youtube.com/channel/ch_1"
    assert analysis.videos[0].daily_views > 0
    assert analysis.unique_channels_count == 2
    assert analysis.validation_summary is not None
    assert analysis.validation_summary.unique_channels_count == 2
    assert analysis.title_analysis is not None
    assert analysis.title_analysis.target_audience == "Beginners"
    assert len(analysis.title_analysis.suggested_title_ideas) > 0
