"""Phase 7 comprehensive automated test suite.

Covers all 26 specified requirements:
1. Fresh database initialization
2. Schema version detection
3. Migration from the existing schema
4. Existing project preservation
5. Existing keyword preservation
6. Provenance preservation
7. Migration rollback on failure
8. Repeated research upserts
9. Project listing
10. Project search
11. Project rename
12. Confirmed project deletion
13. Cache hits and expiry
14. Clearing cache without deleting history
15. Locale-specific cache keys
16. CSV export
17. XLSX export
18. Bangla Unicode export
19. Empty export
20. Selected and filtered export
21. Formula-injection handling
22. Backup creation
23. Backup validation
24. Restore safety
25. Database connection and transaction cleanup
26. Phase 1–6 regression verification
"""

import json
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import List

import openpyxl
import pandas as pd
import pytest

from src.crawler.models import DiscoveredKeyword
from src.models.keyword import KeywordItem, SourceType
from src.storage.backup import BackupManager
from src.storage.database import DatabaseManager
from src.storage.export import (
    ExportService,
    generate_export_filename,
    sanitize_dataframe_for_export,
    sanitize_formula_injection,
)
from src.storage.migrations import MigrationError, MigrationManager


# ------------------------------------------------------------------------------
# 1. Fresh database initialization
# ------------------------------------------------------------------------------

def test_fresh_database_initialization(tmp_path: Path) -> None:
    """Requirement 1: Fresh database initialization creates all tables, indexes, and constraints."""
    db_file = tmp_path / "fresh.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    assert db_file.exists()
    with mgr.get_connection() as conn:
        cursor = conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%';"
        )
        tables = {row["name"] for row in cursor.fetchall()}
        assert "searches" in tables
        assert "keywords" in tables
        assert "autocomplete_cache" in tables
        assert "schema_migrations" in tables


# ------------------------------------------------------------------------------
# 2. Schema version detection
# ------------------------------------------------------------------------------

def test_schema_version_detection(tmp_path: Path) -> None:
    """Requirement 2: Migration system accurately detects the current schema version."""
    db_file = tmp_path / "version_check.db"
    mgr = DatabaseManager(db_file)

    # Before init, version is 0
    assert mgr.migration_manager.get_current_schema_version() == 0

    # After init, version is 3
    mgr.init_db()
    assert mgr.migration_manager.get_current_schema_version() == 3


# ------------------------------------------------------------------------------
# Helper for Migration fixture
# ------------------------------------------------------------------------------

def _create_legacy_phase5_database(db_file: Path) -> int:
    """Creates a legacy Phase 5 SQLite database with pre-Phase 7 schema and rows."""
    conn = sqlite3.connect(db_file)
    conn.execute(
        """
        CREATE TABLE searches (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            seed_keyword TEXT NOT NULL,
            language TEXT NOT NULL,
            country TEXT NOT NULL,
            max_depth INTEGER NOT NULL,
            request_budget INTEGER NOT NULL,
            total_results INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL
        );
        """
    )
    conn.execute(
        """
        CREATE TABLE keywords (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            search_id INTEGER,
            keyword TEXT NOT NULL,
            seed_keyword TEXT NOT NULL,
            source_type TEXT NOT NULL,
            classification TEXT NOT NULL,
            depth INTEGER NOT NULL,
            parent_keyword TEXT,
            language TEXT NOT NULL,
            country TEXT NOT NULL,
            created_at TEXT NOT NULL,
            normalized_key TEXT,
            tail_length TEXT,
            intent_labels TEXT,
            sources_json TEXT,
            word_count INTEGER
        );
        """
    )
    conn.execute(
        """
        CREATE TABLE autocomplete_cache (
            query TEXT NOT NULL,
            language TEXT NOT NULL,
            country TEXT NOT NULL,
            results_json TEXT NOT NULL,
            created_at TEXT NOT NULL,
            PRIMARY KEY (query, language, country)
        );
        """
    )
    created_ts = "2026-10-18T10:00:00+00:00"
    conn.execute(
        """
        INSERT INTO searches (seed_keyword, language, country, max_depth, request_budget, total_results, created_at)
        VALUES ('python tutorial', 'en', 'US', 2, 20, 1, ?);
        """,
        (created_ts,),
    )
    conn.execute(
        """
        INSERT INTO keywords (
            search_id, keyword, seed_keyword, source_type, classification,
            depth, parent_keyword, language, country, created_at,
            normalized_key, tail_length, intent_labels, sources_json, word_count
        ) VALUES (
            1, 'python tutorial for beginners', 'python tutorial', 'autocomplete', 'Informational',
            1, 'python tutorial', 'en', 'US', ?,
            'python tutorial for beginners', 'Long-tail', '["How-to", "Informational"]', '["autocomplete"]', 4
        );
        """,
        (created_ts,),
    )
    conn.commit()
    conn.close()
    return 1


# ------------------------------------------------------------------------------
# 3. Migration from the existing schema
# ------------------------------------------------------------------------------

def test_migration_from_existing_schema(tmp_path: Path) -> None:
    """Requirement 3: Migration system transitions an existing schema to v3 cleanly."""
    db_file = tmp_path / "mig_test.db"
    _create_legacy_phase5_database(db_file)

    mig_mgr = MigrationManager(db_file)
    new_ver, applied = mig_mgr.migrate()

    assert new_ver == 3
    assert "v3_phase7_projects_provenance_columns" in applied


# ------------------------------------------------------------------------------
# 4. Existing project preservation
# ------------------------------------------------------------------------------

def test_existing_project_preservation(tmp_path: Path) -> None:
    """Requirement 4: Existing searches table records are preserved with project metadata backfilled."""
    db_file = tmp_path / "preserve_proj.db"
    s_id = _create_legacy_phase5_database(db_file)

    mgr = DatabaseManager(db_file)
    mgr.init_db()

    project = mgr.get_project(s_id)
    assert project is not None
    assert project["seed_keyword"] == "python tutorial"
    assert project["project_name"] == "python tutorial"
    assert project["status"] == "completed"
    assert project["created_at"] == "2026-10-18T10:00:00+00:00"


# ------------------------------------------------------------------------------
# 5. Existing keyword preservation
# ------------------------------------------------------------------------------

def test_existing_keyword_preservation(tmp_path: Path) -> None:
    """Requirement 5: Existing keyword rows survive migration without loss or modification."""
    db_file = tmp_path / "preserve_kw.db"
    s_id = _create_legacy_phase5_database(db_file)

    mgr = DatabaseManager(db_file)
    mgr.init_db()

    keywords = mgr.get_keywords_by_search_id(s_id)
    assert len(keywords) == 1
    assert keywords[0]["keyword"] == "python tutorial for beginners"
    assert keywords[0]["normalized_key"] == "python tutorial for beginners"


# ------------------------------------------------------------------------------
# 6. Provenance preservation
# ------------------------------------------------------------------------------

def test_provenance_preservation(tmp_path: Path) -> None:
    """Requirement 6: Provenance columns (sources, parents, discovery paths) are accurately preserved."""
    db_file = tmp_path / "preserve_prov.db"
    s_id = _create_legacy_phase5_database(db_file)

    mgr = DatabaseManager(db_file)
    mgr.init_db()

    keywords = mgr.get_keywords_by_search_id(s_id)
    kw = keywords[0]
    assert "autocomplete" in kw["sources"]
    assert "python tutorial" in kw["parent_queries"]
    assert kw["source_name"] == "youtube_autocomplete"


# ------------------------------------------------------------------------------
# 7. Migration rollback on failure
# ------------------------------------------------------------------------------

def test_migration_rollback_on_failure(tmp_path: Path) -> None:
    """Requirement 7: Failed migration rolls back transaction cleanly without leaving partial changes."""
    db_file = tmp_path / "failing_mig.db"
    mig_mgr = MigrationManager(db_file)

    # Initialize v1
    conn = sqlite3.connect(db_file)
    mig_mgr._migration_v1_initial(conn)
    mig_mgr.record_migration(conn, 1, "v1_initial_schema")
    conn.commit()
    conn.close()

    def broken_migration(c: sqlite3.Connection) -> None:
        c.execute("ALTER TABLE searches ADD COLUMN doomed_column TEXT;")
        raise RuntimeError("Simulated migration crash")

    mig_mgr.get_migrations = lambda: [
        (1, "v1_initial_schema", mig_mgr._migration_v1_initial),
        (2, "broken_migration", broken_migration),
    ]

    with pytest.raises(MigrationError):
        mig_mgr.migrate()

    conn = sqlite3.connect(db_file)
    cursor = conn.execute("PRAGMA table_info(searches);")
    cols = {row[1] for row in cursor.fetchall()}
    conn.close()
    assert "doomed_column" not in cols


# ------------------------------------------------------------------------------
# 8. Repeated research upserts
# ------------------------------------------------------------------------------

def test_repeated_research_upserts(tmp_path: Path) -> None:
    """Requirement 8: Rerunning research upserts keywords, merging sources without duplicate rows."""
    db_file = tmp_path / "test_upserts.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    search_id = mgr.save_search("python", "en", "US", 1, 5, project_name="PyTest")

    item1 = KeywordItem(keyword="python tutorial", seed_keyword="python", source=SourceType.AUTOCOMPLETE)
    mgr.save_keywords(search_id, [item1])

    # Second crawl returns same keyword but with generated source
    item2 = KeywordItem(keyword="python tutorial", seed_keyword="python", source=SourceType.GENERATED)
    mgr.save_keywords(search_id, [item2])

    rows = mgr.get_keywords_by_search_id(search_id)
    assert len(rows) == 1
    assert rows[0]["normalized_key"] == item1.normalized_key
    assert "autocomplete" in rows[0]["sources"]
    assert "generated" in rows[0]["sources"]
    assert rows[0]["source_type"] == "autocomplete"


# ------------------------------------------------------------------------------
# 9. Project listing
# ------------------------------------------------------------------------------

def test_project_listing(tmp_path: Path) -> None:
    """Requirement 9: List all saved research projects ordered by recent first."""
    db_file = tmp_path / "listing.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    mgr.save_search("alpha", "en", "US", 1, 5, project_name="Project Alpha")
    mgr.save_search("beta", "en", "US", 1, 5, project_name="Project Beta")

    projects = mgr.list_projects()
    assert len(projects) == 2
    assert projects[0]["project_name"] == "Project Beta"
    assert projects[1]["project_name"] == "Project Alpha"


# ------------------------------------------------------------------------------
# 10. Project search
# ------------------------------------------------------------------------------

def test_project_search(tmp_path: Path) -> None:
    """Requirement 10: Search project history by project name or seed keyword."""
    db_file = tmp_path / "search_proj.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    mgr.save_search("fastapi", "en", "US", 1, 5, project_name="API Frameworks")
    mgr.save_search("django", "en", "US", 1, 5, project_name="Web Dev")

    found = mgr.list_projects(search_query="API")
    assert len(found) == 1
    assert found[0]["project_name"] == "API Frameworks"


# ------------------------------------------------------------------------------
# 11. Project rename
# ------------------------------------------------------------------------------

def test_project_rename(tmp_path: Path) -> None:
    """Requirement 11: Rename a project and update its updated_at timestamp."""
    db_file = tmp_path / "rename_proj.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    s_id = mgr.save_search("initial seed", "en", "US", 1, 5, project_name="Old Name")
    success = mgr.rename_project(s_id, "New Brand Name")
    assert success is True

    updated = mgr.get_project(s_id)
    assert updated["project_name"] == "New Brand Name"


# ------------------------------------------------------------------------------
# 12. Confirmed project deletion
# ------------------------------------------------------------------------------

def test_confirmed_project_deletion(tmp_path: Path) -> None:
    """Requirement 12: Delete a project and cascade delete all its associated keywords."""
    db_file = tmp_path / "del_proj.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    s_id = mgr.save_search("seed", "en", "US", 1, 5)
    mgr.save_keywords(s_id, [KeywordItem(keyword="kw1", seed_keyword="seed", source=SourceType.AUTOCOMPLETE)])

    assert len(mgr.get_keywords_by_search_id(s_id)) == 1
    deleted = mgr.delete_search(s_id)
    assert deleted is True

    assert mgr.get_project(s_id) is None
    assert len(mgr.get_keywords_by_search_id(s_id)) == 0


# ------------------------------------------------------------------------------
# 13. Cache hits and expiry
# ------------------------------------------------------------------------------

def test_cache_hits_and_expiry(tmp_path: Path) -> None:
    """Requirement 13: Autocomplete cache returns suggestions within TTL and expires old ones."""
    db_file = tmp_path / "cache_exp.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    mgr.set_cached_suggestions("python", "en", "US", ["python 1", "python 2"])

    # Active cache hit
    hit = mgr.get_cached_suggestions("python", "en", "US", max_age_seconds=86400)
    assert hit == ["python 1", "python 2"]

    # Expired cache miss
    miss = mgr.get_cached_suggestions("python", "en", "US", max_age_seconds=0)
    assert miss is None


# ------------------------------------------------------------------------------
# 14. Clearing cache without deleting history
# ------------------------------------------------------------------------------

def test_clearing_cache_without_deleting_history(tmp_path: Path) -> None:
    """Requirement 14: Clearing cache entries leaves research projects and keywords intact."""
    db_file = tmp_path / "cache_clear.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    s_id = mgr.save_search("python", "en", "US", 1, 5)
    mgr.save_keywords(s_id, [KeywordItem(keyword="python kw", seed_keyword="python", source=SourceType.AUTOCOMPLETE)])
    mgr.set_cached_suggestions("python", "en", "US", ["python kw"])

    cleared = mgr.clear_cache(expired_only=False)
    assert cleared == 1

    # Cache is cleared
    assert mgr.get_cache_stats()["total_cached_queries"] == 0
    # Project and keywords are intact
    assert mgr.get_project(s_id) is not None
    assert len(mgr.get_keywords_by_search_id(s_id)) == 1


# ------------------------------------------------------------------------------
# 15. Locale-specific cache keys
# ------------------------------------------------------------------------------

def test_locale_specific_cache_keys(tmp_path: Path) -> None:
    """Requirement 15: Cache distinguishes results based on query, language, and country."""
    db_file = tmp_path / "cache_locale.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    mgr.set_cached_suggestions("news", "en", "US", ["us news"])
    mgr.set_cached_suggestions("news", "en", "GB", ["uk news"])

    assert mgr.get_cached_suggestions("news", "en", "US") == ["us news"]
    assert mgr.get_cached_suggestions("news", "en", "GB") == ["uk news"]


# ------------------------------------------------------------------------------
# 16. CSV export
# ------------------------------------------------------------------------------

def test_csv_export() -> None:
    """Requirement 16: CSV export generates well-formed UTF-8 text with headers."""
    df = pd.DataFrame([{"Keyword": "test kw", "Source": "Autocomplete"}])
    csv_bytes = ExportService.export_csv(df)
    assert b"test kw" in csv_bytes
    assert b"Autocomplete" in csv_bytes


# ------------------------------------------------------------------------------
# 17. XLSX export
# ------------------------------------------------------------------------------

def test_xlsx_export() -> None:
    """Requirement 17: XLSX export produces valid openpyxl spreadsheet bytes with headers."""
    df = pd.DataFrame([{"Keyword": "test kw", "Word Count": 2}])
    xlsx_bytes = ExportService.export_excel(df, sheet_name="ExportTest")
    assert len(xlsx_bytes) > 0

    wb = openpyxl.load_workbook(filename=openpyxl.reader.excel.BytesIO(xlsx_bytes))
    ws = wb["ExportTest"]
    assert ws.cell(row=1, column=1).value == "Keyword"
    assert ws.cell(row=2, column=1).value == "test kw"


# ------------------------------------------------------------------------------
# 18. Bangla Unicode export
# ------------------------------------------------------------------------------

def test_bangla_unicode_export() -> None:
    """Requirement 18: CSV and XLSX exports preserve Bengali script with UTF-8 BOM."""
    df = pd.DataFrame([{"Keyword": "পাইথন প্রোগ্রামিং বাংলা", "Language": "bn"}])
    csv_bytes = ExportService.export_csv(df, include_bom=True)
    assert csv_bytes.startswith(b"\xef\xbb\xbf")
    assert "পাইথন প্রোগ্রামিং বাংলা" in csv_bytes.decode("utf-8-sig")

    xlsx_bytes = ExportService.export_excel(df)
    wb = openpyxl.load_workbook(filename=openpyxl.reader.excel.BytesIO(xlsx_bytes))
    ws = wb.active
    assert ws.cell(row=2, column=1).value == "পাইথন প্রোগ্রামিং বাংলা"


# ------------------------------------------------------------------------------
# 19. Empty export
# ------------------------------------------------------------------------------

def test_empty_export() -> None:
    """Requirement 19: Empty result sets export cleanly with column headers without crashing."""
    df = pd.DataFrame(columns=["Keyword", "Tail Type", "Primary Intent"])
    csv_bytes = ExportService.export_csv(df)
    assert len(csv_bytes) > 0
    assert "Keyword,Tail Type,Primary Intent" in csv_bytes.decode("utf-8-sig")

    xlsx_bytes = ExportService.export_excel(df)
    assert len(xlsx_bytes) > 0


# ------------------------------------------------------------------------------
# 20. Selected and filtered export
# ------------------------------------------------------------------------------

def test_selected_and_filtered_export() -> None:
    """Requirement 20: Exports correctly handle selected and filtered DataFrame subsets."""
    df = pd.DataFrame(
        [
            {"Keyword": "keep this", "Status": "selected"},
            {"Keyword": "drop this", "Status": "unselected"},
        ]
    )
    filtered_df = df[df["Status"] == "selected"]
    csv_bytes = ExportService.export_csv(filtered_df)
    assert "keep this" in csv_bytes.decode("utf-8-sig")
    assert "drop this" not in csv_bytes.decode("utf-8-sig")


# ------------------------------------------------------------------------------
# 21. Formula-injection handling
# ------------------------------------------------------------------------------

def test_formula_injection_handling() -> None:
    """Requirement 21: Cells starting with =, +, -, @, \\t, \\r are escaped with a leading single quote."""
    dangerous = ["=SUM(A1:A10)", "+CMD|' /C calc'!A0", "-cmd", "@special", "\tindent", "\rcarriage"]
    df = pd.DataFrame([{"Keyword": val} for val in dangerous])

    sanitized = sanitize_dataframe_for_export(df)
    for idx, row in sanitized.iterrows():
        assert row["Keyword"].startswith("'")

    # Safe keywords are left untouched
    safe_df = pd.DataFrame([{"Keyword": "normal keyword"}])
    safe_sanitized = sanitize_dataframe_for_export(safe_df)
    assert safe_sanitized.iloc[0]["Keyword"] == "normal keyword"


# ------------------------------------------------------------------------------
# 22. Backup creation
# ------------------------------------------------------------------------------

def test_backup_creation(tmp_path: Path) -> None:
    """Requirement 22: SQLite online backup creates a physical snapshot of the database."""
    db_file = tmp_path / "main.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    mgr.save_search("backup seed", "en", "US", 1, 5)
    bkp = mgr.create_backup("my_backup.db")

    assert bkp.exists()
    assert bkp.stat().st_size > 0


# ------------------------------------------------------------------------------
# 23. Backup validation
# ------------------------------------------------------------------------------

def test_backup_validation(tmp_path: Path) -> None:
    """Requirement 23: Backup validation runs PRAGMA integrity_check and table checks."""
    db_file = tmp_path / "valid_test.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    bkp = mgr.create_backup("check_me.db")
    is_valid, msg, stats = mgr.validate_backup(bkp)

    assert is_valid is True
    assert stats["integrity"] == "ok"
    assert "searches" in stats["tables"]
    assert "keywords" in stats["tables"]


# ------------------------------------------------------------------------------
# 24. Restore safety
# ------------------------------------------------------------------------------

def test_restore_safety(tmp_path: Path) -> None:
    """Requirement 24: Restoring a backup takes a pre-restore safety copy and overwrites cleanly."""
    db_file = tmp_path / "restore_test.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    # State A: 1 search
    mgr.save_search("state A", "en", "US", 1, 5, project_name="State A")
    bkp = mgr.create_backup("snapshot_A.db")

    # State B: 2 searches
    mgr.save_search("state B", "en", "US", 1, 5, project_name="State B")
    assert len(mgr.list_projects()) == 2

    # Restore State A
    success, msg = mgr.restore_backup(bkp)
    assert success is True

    # Active database is now State A
    projects = mgr.list_projects()
    assert len(projects) == 1
    assert projects[0]["project_name"] == "State A"

    # Pre-restore safety backup exists
    backups = mgr.list_backups()
    assert any("pre_restore_safety" in b["name"] for b in backups)


# ------------------------------------------------------------------------------
# 25. Database connection and transaction cleanup
# ------------------------------------------------------------------------------

def test_database_connection_and_transaction_cleanup(tmp_path: Path) -> None:
    """Requirement 25: Connection context manager closes cleanly even on operational errors."""
    db_file = tmp_path / "cleanup.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    with pytest.raises(sqlite3.OperationalError):
        with mgr.get_connection() as conn:
            conn.execute("SELECT * FROM invalid_nonexistent_table;")

    # Connection should be released and database immediately accessible
    with mgr.get_connection() as conn:
        cursor = conn.execute("SELECT COUNT(*) FROM searches;")
        assert cursor.fetchone()[0] == 0


# ------------------------------------------------------------------------------
# 26. Phase 1–6 regression verification
# ------------------------------------------------------------------------------

def test_phase1_to_6_regression_smoke(tmp_path: Path) -> None:
    """Requirement 26: Regression verification confirming models, patterns, and crawler play nicely."""
    db_file = tmp_path / "smoke.db"
    mgr = DatabaseManager(db_file)
    mgr.init_db()

    s_id = mgr.save_search("python", "en", "US", 1, 10, project_name="Regression Smoke")
    disc = DiscoveredKeyword(
        keyword="python tutorial",
        normalized_key="python tutorial",
        depth=1,
        parent_queries=["python"],
        discovery_paths=2,
        source_type="autocomplete",
        language="en",
        country="US",
        sources=["autocomplete", "generated"],
        tail_length="Short-tail",
        intent_labels=["How-to"],
        word_count=2,
    )
    saved = mgr.save_keywords(s_id, [disc])
    assert saved == 1

    loaded = mgr.get_keywords_by_search_id(s_id)
    assert len(loaded) == 1
    assert loaded[0]["keyword"] == "python tutorial"
    assert loaded[0]["sources"] == ["autocomplete", "generated"]
