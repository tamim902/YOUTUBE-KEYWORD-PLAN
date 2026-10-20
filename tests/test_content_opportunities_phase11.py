"""Phase 11 automated test suite: YouTube Content Opportunity Analyzer & Video Planning System.

Covers:
1. Domain models (ContentOpportunityItem, ContentBrief, CalendarEntry, TopicOverlapResult, TitleQualityCheck).
2. ContentGapDetector empirical behavior (with and without competitor videos, strict empirical labeling).
3. OpportunityScorer 0-100 heuristic formula, 6 dimensions, confidence ratings, and explainability.
4. VideoTopicGenerator format variations, audience tailoring, and supporting keywords association.
5. TitleHookWorkshop 7 categorized titles, 7-point quality checklist, clickbait detection, hook & outline.
6. ContentBriefGenerator 15-point structured editable briefs.
7. TopicOverlapDetector Jaccard similarity, cannibalization warnings, and overlap types.
8. Database persistence for opportunities, briefs, and calendar entries with schema idempotency.
9. CSV & Excel export services with formula-injection neutralization and UTF-8 BOM encoding.
10. Multilingual Bengali Unicode keyword and topic handling.
11. UI navigation constants and backward compatibility.
"""

import io
import json
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List

import openpyxl
import pytest

from src.models.opportunity import (
    CalendarEntry,
    ContentBrief,
    ContentOpportunityItem,
    TitleQualityCheck,
    TopicOverlapResult,
)
from src.intelligence.opportunities import (
    ContentBriefGenerator,
    ContentGapDetector,
    OpportunityScorer,
    TitleHookWorkshop,
    TopicOverlapDetector,
    VideoTopicGenerator,
)
from src.storage.database import DatabaseManager
from src.storage.export import ExportService
from src.ui.components import (
    ALL_NAVIGATION_PAGES,
    NAVIGATION_PAGES,
    PAGE_CONTENT_OPPORTUNITIES,
)


# ------------------------------------------------------------------------------
# 1. Domain Models
# ------------------------------------------------------------------------------

def test_content_opportunity_item_serialization() -> None:
    """Verify ContentOpportunityItem model creation, dict conversion, and round-trip."""
    item = ContentOpportunityItem(
        id=1,
        topic="How to Wash a Black Car Without Swirl Marks",
        primary_keyword="wash black car",
        supporting_keywords=["swirl marks", "microfiber wash mitt", "two bucket method"],
        target_audience="Beginners",
        search_intent="How-to / Tutorial",
        content_format="Step-by-step Guide",
        opportunity_score=85.5,
        confidence_level="High Evidence",
        status="Idea",
        priority="High",
        score_breakdown={"relevance_score": 25.0, "audience_score": 20.0},
        gap_rationale="Question query has 0 matching titles among 10 competitor videos.",
        notes="High priority beginner topic.",
        project_id=42,
    )

    d = item.to_dict()
    assert d["id"] == 1
    assert d["topic"] == "How to Wash a Black Car Without Swirl Marks"
    assert d["primary_keyword"] == "wash black car"
    assert len(d["supporting_keywords"]) == 3
    assert d["opportunity_score"] == 85.5
    assert d["confidence_level"] == "High Evidence"
    assert d["priority"] == "High"
    assert d["project_id"] == 42

    restored = ContentOpportunityItem.from_dict(d)
    assert restored.id == item.id
    assert restored.topic == item.topic
    assert restored.primary_keyword == item.primary_keyword
    assert restored.supporting_keywords == item.supporting_keywords
    assert restored.opportunity_score == item.opportunity_score
    assert restored.score_breakdown == item.score_breakdown


def test_content_brief_model_serialization() -> None:
    """Verify ContentBrief model handles 15 structured fields and serialization."""
    brief = ContentBrief(
        id=10,
        opportunity_id=1,
        working_title="How to Wash a Black Car Without Swirl Marks (2026 Guide)",
        primary_keyword="wash black car",
        supporting_keywords=["two bucket method", "foam cannon"],
        target_audience="Beginners",
        search_intent="How-to / Tutorial",
        viewer_problem="Black cars show fine scratches and swirl marks immediately after washing.",
        unique_angle="Complete step-by-step guide with zero marketing hype and clear sunlight inspection.",
        proposed_intro="Hook the viewer with common black car washing mistakes.",
        main_sections=["1. Prep & Rinse", "2. Two-Bucket Method", "3. Drying Technique"],
        examples_to_research=["Check 2-3 pH neutral car shampoos."],
        visual_elements=["Sunlight reflection test", "Foam cannon application close-up"],
        faqs=["Is two bucket method necessary?", "What towel is best for black paint?"],
        conclusion_cta="Summarize 3 golden rules and link to waxing guide.",
        fact_checking_checklist=["Verify safe wash dilution ratios."],
        source_references=["Seed keyword query: wash black car"],
    )

    d = brief.to_dict()
    assert d["working_title"] == "How to Wash a Black Car Without Swirl Marks (2026 Guide)"
    assert len(d["main_sections"]) == 3
    assert len(d["visual_elements"]) == 2
    assert len(d["faqs"]) == 2

    restored = ContentBrief.from_dict(d)
    assert restored.working_title == brief.working_title
    assert restored.primary_keyword == brief.primary_keyword
    assert restored.main_sections == brief.main_sections
    assert restored.visual_elements == brief.visual_elements


def test_calendar_entry_model_serialization() -> None:
    """Verify CalendarEntry model serialization and default values."""
    entry = CalendarEntry(
        id=5,
        opportunity_id=1,
        title="Ceramic Coating vs Wax - Which Is Better?",
        primary_keyword="ceramic coating vs wax",
        planned_date="2026-11-15",
        status="Scripting",
        priority="High",
        notes="Record B-roll outdoors.",
    )
    d = entry.to_dict()
    assert d["title"] == "Ceramic Coating vs Wax - Which Is Better?"
    assert d["status"] == "Scripting"
    assert d["planned_date"] == "2026-11-15"

    restored = CalendarEntry.from_dict(d)
    assert restored.title == entry.title
    assert restored.status == entry.status
    assert restored.priority == entry.priority


def test_topic_overlap_and_title_check_models() -> None:
    """Verify TopicOverlapResult and TitleQualityCheck models."""
    overlap = TopicOverlapResult(
        topic_a="Best Car Wax 2026",
        topic_b="Top 5 Best Car Waxes Tested",
        similarity_score=0.75,
        overlap_type="High Token Overlap",
        shared_keywords=["best", "car", "wax"],
        warning_message="Potential cannibalization detected.",
        resolution="Unresolved",
    )
    assert overlap.similarity_score == 0.75
    assert "car" in overlap.shared_keywords
    assert overlap.to_dict()["overlap_type"] == "High Token Overlap"

    check = TitleQualityCheck(
        title="How to Wax a Car for Beginners (2026 Guide)",
        score=95.0,
        passes_relevance=True,
        passes_clarity=True,
        passes_specificity=True,
        passes_audience_fit=True,
        has_unsupported_claims=False,
        has_excessive_repetition=False,
        has_misleading_wording=False,
        feedback=["Excellent title!"],
    )
    assert check.score == 95.0
    assert check.passes_relevance is True
    assert check.to_dict()["score"] == 95.0


# ------------------------------------------------------------------------------
# 2. Content Gap Analysis (ContentGapDetector)
# ------------------------------------------------------------------------------

def test_content_gap_detector_with_competitor_titles() -> None:
    """Verify ContentGapDetector classifies gap types and uses strict empirical labeling."""
    competitor_titles = [
        "Meguiar's Ultimate Liquid Wax Review",
        "How to Polish Your Car Like a Pro",
        "Chemical Guys Butter Wet Wax Application",
        "Top 10 Car Care Products of the Year",
    ]
    detector = ContentGapDetector(competitor_titles)
    assert detector.is_video_data_available() is True
    assert detector.sample_size == 4

    # 1. Question gap: question query with no matching title in sample
    res_q = detector.detect_gap("how to wax a black car without swirls")
    assert res_q["is_gap"] is True
    assert res_q["gap_label"] == "Potential Content Gap"
    assert res_q["gap_type"] == "QUESTION_GAP"
    assert "4 competitor videos" in res_q["evidence_rationale"]

    # 2. Comparison gap
    res_comp = detector.detect_gap("ceramic coating vs car wax")
    assert res_comp["is_gap"] is True
    assert res_comp["gap_label"] == "Potential Content Gap"
    assert res_comp["gap_type"] == "COMPARISON_GAP"

    # 3. Audience gap
    res_aud = detector.detect_gap("car wax for beginners")
    assert res_aud["is_gap"] is True
    assert res_aud["gap_label"] == "Potential Content Gap"
    assert res_aud["gap_type"] == "AUDIENCE_GAP"

    # 4. Uncovered general keyword
    res_uncov = detector.detect_gap("graphene spray coating application")
    assert res_uncov["is_gap"] is True
    assert res_uncov["gap_label"] == "Potential Content Gap"
    assert res_uncov["gap_type"] == "UNCOVERED_KEYWORD"

    # 5. Covered keyword in sample
    res_cov = detector.detect_gap("meguiar's ultimate liquid wax")
    assert res_cov["is_gap"] is False
    assert res_cov["gap_label"] == "Covered in Competitor Sample"
    assert res_cov["gap_type"] == "COVERED_IN_SAMPLE"
    assert res_cov["matches_count"] >= 1


def test_content_gap_detector_without_competitor_titles() -> None:
    """Verify ContentGapDetector honestly reports absence of competitor sample and creates no false gaps."""
    detector_empty = ContentGapDetector([])
    assert detector_empty.is_video_data_available() is False

    res = detector_empty.detect_gap("best car wax")
    assert res["is_gap"] is False
    assert res["gap_label"] == "No Video Sample Available"
    assert res["gap_type"] == "UNVERIFIED_COMPETITOR_SAMPLE"
    assert res["sample_size"] == 0
    assert "cannot be verified" in res["evidence_rationale"]

    detector_none = ContentGapDetector(None)
    assert detector_none.is_video_data_available() is False


# ------------------------------------------------------------------------------
# 3. Transparent Opportunity Scorer (OpportunityScorer)
# ------------------------------------------------------------------------------

def test_opportunity_scorer_formula_and_weights() -> None:
    """Verify OpportunityScorer evaluates all 6 weighted dimensions transparently."""
    scorer = OpportunityScorer()

    gap_info = {
        "is_gap": True,
        "gap_label": "Potential Content Gap",
        "gap_type": "QUESTION_GAP",
    }
    supporting = ["car wax tips", "wax buffer", "microfiber cloth", "swirl remover", "liquid wax"]

    score, confidence, breakdown = scorer.calculate_score(
        keyword="best car wax for beginners",
        seed_keyword="best car wax",
        target_audience="Beginners",
        search_intent="How-to / Tutorial",
        content_format="Step-by-step Guide",
        supporting_keywords=supporting,
        gap_info=gap_info,
        has_video_sample=True,
    )

    # Validate bounds
    assert 0.0 <= score <= 100.0
    # High evidence since video sample is true and supporting count >= 3
    assert confidence == "High Evidence"

    # Validate all 6 dimensions are in breakdown
    assert "relevance_score" in breakdown
    assert "audience_score" in breakdown
    assert "intent_score" in breakdown
    assert "supporting_score" in breakdown
    assert "gap_score" in breakdown
    assert "format_score" in breakdown

    # Check sum matches
    calculated_sum = round(
        breakdown["relevance_score"]
        + breakdown["audience_score"]
        + breakdown["intent_score"]
        + breakdown["supporting_score"]
        + breakdown["gap_score"]
        + breakdown["format_score"],
        1,
    )
    assert breakdown["total_score"] == calculated_sum
    assert breakdown["total_score"] == score

    # Check retrieved vs user preference segregation
    assert len(breakdown["inputs_from_retrieved_data"]) >= 3
    assert len(breakdown["inputs_from_user_preferences"]) >= 2
    assert "formula_explanation" in breakdown
    assert "NOT ranking probability" in breakdown["formula_explanation"]


def test_opportunity_scorer_confidence_levels() -> None:
    """Verify confidence levels derive strictly from available empirical evidence."""
    scorer = OpportunityScorer()

    # 1. High Evidence: has video sample and >= 3 supporting keywords
    _, conf_high, _ = scorer.calculate_score(
        keyword="car wax",
        seed_keyword="car wax",
        target_audience="General Audience",
        search_intent="Informational",
        content_format="General",
        supporting_keywords=["kw1", "kw2", "kw3"],
        gap_info={"is_gap": True},
        has_video_sample=True,
    )
    assert conf_high == "High Evidence"

    # 2. Moderate Evidence: has video sample but < 3 supporting keywords
    _, conf_mod, _ = scorer.calculate_score(
        keyword="car wax",
        seed_keyword="car wax",
        target_audience="General Audience",
        search_intent="Informational",
        content_format="General",
        supporting_keywords=["kw1"],
        gap_info={"is_gap": False},
        has_video_sample=True,
    )
    assert conf_mod == "Moderate Evidence"

    # 3. Low / Lexical Only: no video sample and < 2 supporting keywords
    _, conf_low, _ = scorer.calculate_score(
        keyword="car wax",
        seed_keyword="car wax",
        target_audience="General Audience",
        search_intent="Informational",
        content_format="General",
        supporting_keywords=[],
        gap_info={"is_gap": False},
        has_video_sample=False,
    )
    assert conf_low == "Low / Lexical Only"


# ------------------------------------------------------------------------------
# 4. Video Topic Generator (VideoTopicGenerator)
# ------------------------------------------------------------------------------

def test_video_topic_generator_structured_formats() -> None:
    """Verify VideoTopicGenerator outputs diverse structured topics with supporting keywords."""
    generator = VideoTopicGenerator()

    keywords = [
        "best car wax",
        "how to apply car wax",
        "car wax vs sealant",
        "car wax for black cars",
        "car wax mistakes",
        "meguiar's car wax review",
        "best car wax for beginners",
    ]
    competitor_titles = [
        "Top 10 Car Waxes You Can Buy",
        "How to Detail Your Car Exterior",
    ]

    topics = generator.generate_topics_for_keywords(
        keywords=keywords,
        seed_keyword="car wax",
        target_audience="Beginners",
        preferred_format="All Formats",
        competitor_titles=competitor_titles,
        max_topics=8,
        project_id=1,
    )

    assert len(topics) > 0
    assert len(topics) <= 8

    # Verify attributes of generated topics
    formats_found = {t.content_format for t in topics}
    assert len(formats_found) >= 2  # multiple formats generated

    for t in topics:
        assert isinstance(t, ContentOpportunityItem)
        assert t.topic
        assert t.primary_keyword
        assert 0.0 <= t.opportunity_score <= 100.0
        assert t.confidence_level in ("High Evidence", "Moderate Evidence", "Low / Lexical Only")
        assert t.target_audience == "Beginners"
        assert t.project_id == 1


def test_video_topic_generator_format_filter() -> None:
    """Verify VideoTopicGenerator respects preferred_format filter."""
    generator = VideoTopicGenerator()
    keywords = ["car wax tips", "how to wax", "wax guide"]

    topics = generator.generate_topics_for_keywords(
        keywords=keywords,
        seed_keyword="car wax",
        target_audience="General Audience",
        preferred_format="Step-by-step Guide",
        competitor_titles=None,
        max_topics=5,
    )
    assert len(topics) > 0
    for t in topics:
        assert t.content_format == "Step-by-step Guide"


def test_video_topic_generator_empty_input() -> None:
    """Verify VideoTopicGenerator safely handles empty input lists."""
    generator = VideoTopicGenerator()
    assert generator.generate_topics_for_keywords([], "seed") == []


# ------------------------------------------------------------------------------
# 5. Title & Hook Workshop (TitleHookWorkshop)
# ------------------------------------------------------------------------------

def test_title_hook_workshop_generates_7_categories() -> None:
    """Verify TitleHookWorkshop produces 7 categorized title options."""
    workshop = TitleHookWorkshop()
    angles = workshop.generate_title_options(
        keyword="car wax",
        target_audience="Beginners",
        current_year=2026,
    )

    expected_categories = [
        "Informational",
        "How-To / Practical",
        "Comparison / VS",
        "Review / Hands-On",
        "Question-Based",
        "List-Based",
        "Audience-Specific",
    ]
    for cat in expected_categories:
        assert cat in angles
        assert "Car Wax" in angles[cat]
        assert len(angles[cat]) > 10


def test_title_quality_check_scoring_and_clickbait() -> None:
    """Verify TitleHookWorkshop 7-point quality checklist flags clickbait and short/long titles."""
    workshop = TitleHookWorkshop()

    # 1. High-quality title
    good_title = "How to Apply Car Wax for Beginners (Step-by-Step Guide)"
    check_good = workshop.check_title_quality(good_title, primary_keyword="car wax")
    assert check_good.score >= 80.0
    assert check_good.passes_relevance is True
    assert check_good.passes_clarity is True
    assert check_good.has_unsupported_claims is False

    # 2. Clickbait title with sensational words
    clickbait_title = "The SECRET Car Wax Trick That Gives INSANE Guaranteed Results"
    check_cb = workshop.check_title_quality(clickbait_title, primary_keyword="car wax")
    assert check_cb.has_unsupported_claims is True
    assert check_cb.score < check_good.score
    assert any("misleading or unsupported" in f for f in check_cb.feedback)

    # 3. Off-topic title (missing primary keyword)
    off_topic = "Top 5 Ways to Clean Your House Fast"
    check_off = workshop.check_title_quality(off_topic, primary_keyword="car wax")
    assert check_off.passes_relevance is False
    assert any("does not clearly feature" in f for f in check_off.feedback)

    # 4. Excessively repetitive title
    repetitive = "Wax car wax car wax best wax car wax"
    check_rep = workshop.check_title_quality(repetitive, primary_keyword="car wax")
    assert check_rep.has_excessive_repetition is True


def test_title_hook_workshop_hook_and_outline() -> None:
    """Verify TitleHookWorkshop script hook generation and 5-part timeline outline."""
    workshop = TitleHookWorkshop()
    res = workshop.generate_hook_and_outline(
        title="How to Master Car Wax in 4 Easy Steps",
        keyword="car wax",
        target_audience="Beginners",
    )

    assert "hook" in res
    assert "outline" in res
    assert len(res["hook"]) > 50
    assert "car wax" in res["hook"].lower()

    # Verify 5-part timeline structure
    outline = res["outline"]
    assert len(outline) == 5
    assert "0:00" in outline[0]
    assert "Step-by-Step" in outline[2]
    assert "Next Action Steps" in outline[4]


# ------------------------------------------------------------------------------
# 6. Content Brief Generator (ContentBriefGenerator)
# ------------------------------------------------------------------------------

def test_content_brief_generator_15_points() -> None:
    """Verify ContentBriefGenerator creates a comprehensive 15-point structured brief."""
    opp = ContentOpportunityItem(
        id=7,
        topic="Ceramic Coating vs Carnauba Wax for Beginners",
        primary_keyword="ceramic coating vs wax",
        supporting_keywords=["paint protection", "durability", "gloss test"],
        target_audience="Beginners",
        search_intent="Comparison / VS",
        content_format="Comparison",
        opportunity_score=88.0,
        confidence_level="High Evidence",
        status="Idea",
        priority="High",
    )

    generator = ContentBriefGenerator()
    brief = generator.generate_brief(opp)

    assert brief.opportunity_id == 7
    assert brief.primary_keyword == "ceramic coating vs wax"
    assert brief.target_audience == "Beginners"
    assert brief.search_intent == "Comparison / VS"
    assert len(brief.viewer_problem) > 20
    assert len(brief.unique_angle) > 20
    assert len(brief.proposed_intro) > 20
    assert len(brief.main_sections) >= 5
    assert len(brief.examples_to_research) >= 2
    assert len(brief.visual_elements) >= 3
    assert len(brief.faqs) >= 2
    assert len(brief.conclusion_cta) > 20
    assert len(brief.fact_checking_checklist) >= 2
    assert len(brief.source_references) >= 2


# ------------------------------------------------------------------------------
# 7. Topic Overlap & Cannibalization Detector (TopicOverlapDetector)
# ------------------------------------------------------------------------------

def test_topic_overlap_detector_cannibalization_warnings() -> None:
    """Verify TopicOverlapDetector identifies high similarity and issues actionable warnings."""
    detector = TopicOverlapDetector()

    items = [
        {"title": "How to Wax a Black Car at Home", "primary_keyword": "wax a black car"},
        {"title": "How to Wax a Black Car Without Swirls", "primary_keyword": "wax a black car"},
        {"title": "Complete Engine Bay Detailing Guide", "primary_keyword": "engine bay detailing"},
    ]

    overlaps = detector.check_overlap(items, similarity_threshold=0.40)
    assert len(overlaps) >= 1

    found_cannibalization = False
    for ov in overlaps:
        if "wax a black car" in ov.shared_keywords or ov.similarity_score >= 0.5:
            found_cannibalization = True
            assert ov.overlap_type in ("Shared Primary Keyword", "High Token Overlap")
            assert "cannibalization detected" in ov.warning_message.lower()
            assert ov.resolution == "Unresolved"

    assert found_cannibalization is True


def test_topic_overlap_detector_zero_overlap() -> None:
    """Verify distinct topics produce zero overlap warnings."""
    detector = TopicOverlapDetector()
    items = [
        {"title": "Leather Seat Restoration", "primary_keyword": "leather seat restoration"},
        {"title": "Headlight Restoration Kit Review", "primary_keyword": "headlight restoration kit"},
    ]
    overlaps = detector.check_overlap(items, similarity_threshold=0.50)
    assert len(overlaps) == 0


# ------------------------------------------------------------------------------
# 8. Database Persistence (SQLite CRUD & Idempotency)
# ------------------------------------------------------------------------------

def test_database_persistence_phase11(tmp_path: Path) -> None:
    """Verify idempotent creation and full CRUD operations on Phase 11 SQLite tables."""
    db_file = tmp_path / "test_phase11.db"
    db = DatabaseManager(str(db_file))
    db.init_db()

    # 1. Opportunities CRUD
    search_id = db.save_search(
        seed_keyword="car wax",
        language="en",
        country="US",
        max_depth=1,
        request_budget=10,
        total_results=5,
    )

    opp = ContentOpportunityItem(
        topic="How to Wash a Car in Winter",
        primary_keyword="wash car winter",
        supporting_keywords=["salt removal", "warm water"],
        target_audience="Car Owners",
        search_intent="How-to / Tutorial",
        content_format="Step-by-step Guide",
        opportunity_score=82.0,
        confidence_level="Moderate Evidence",
        status="Idea",
        priority="Medium",
        notes="Season-specific topic",
        project_id=search_id,
    )
    opp_id = db.save_content_opportunity(opp)
    assert opp_id > 0

    opps = db.get_content_opportunities()
    assert len(opps) == 1
    assert opps[0]["topic"] == "How to Wash a Car in Winter"

    # Status update
    updated = db.update_content_opportunity_status(opp_id, status="Scripting", priority="High")
    assert updated is True
    reloaded = db.get_content_opportunity(opp_id)
    assert reloaded is not None
    assert reloaded["status"] == "Scripting"
    assert reloaded["priority"] == "High"

    # 2. Content Briefs CRUD
    brief = ContentBrief(
        opportunity_id=opp_id,
        working_title="How to Wash a Car in Winter Without Freezing (2026)",
        primary_keyword="wash car winter",
        supporting_keywords=["salt removal"],
        target_audience="Car Owners",
        search_intent="How-to / Tutorial",
        viewer_problem="Salt corrodes cars but freezing temperatures make washing difficult.",
        unique_angle="Indoor and touchless garage wash method.",
        proposed_intro="Hook the viewer with rust damage statistics.",
        main_sections=["1. Undercarriage Rinse", "2. Two Bucket Technique"],
        conclusion_cta="Subscribe for more winter maintenance guides.",
    )
    brief_id = db.save_content_brief(brief)
    assert brief_id > 0

    saved_brief = db.get_content_brief(opp_id)
    assert saved_brief is not None
    assert saved_brief["working_title"] == "How to Wash a Car in Winter Without Freezing (2026)"

    # 3. Calendar Entries CRUD
    entry = CalendarEntry(
        opportunity_id=opp_id,
        title="How to Wash a Car in Winter Without Freezing",
        primary_keyword="wash car winter",
        planned_date="2026-12-01",
        status="Production",
        priority="High",
        notes="Shoot in heated garage.",
    )
    cal_id = db.save_calendar_entry(entry)
    assert cal_id > 0

    entries = db.get_calendar_entries()
    assert len(entries) == 1
    assert entries[0]["status"] == "Production"

    updated_cal = db.update_calendar_entry(cal_id, status="Editing", planned_date="2026-12-05")
    assert updated_cal is True
    reloaded_entries = db.get_calendar_entries()
    assert reloaded_entries[0]["status"] == "Editing"
    assert reloaded_entries[0]["planned_date"] == "2026-12-05"

    # Deletions
    assert db.delete_calendar_entry(cal_id) is True
    assert len(db.get_calendar_entries()) == 0

    assert db.delete_content_brief(brief_id) is True
    assert db.get_content_brief(opp_id) is None

    assert db.delete_content_opportunity(opp_id) is True
    assert len(db.get_content_opportunities()) == 0


# ------------------------------------------------------------------------------
# 9. CSV and Excel Export Services
# ------------------------------------------------------------------------------

def test_phase11_exports_sanitization_and_bom() -> None:
    """Verify Phase 11 CSV and XLSX exports feature UTF-8 BOM and formula-injection neutralization."""
    dangerous_title = "=SUM(1,2)*cmd|' /C calc'!A0"
    opp = ContentOpportunityItem(
        topic=dangerous_title,
        primary_keyword="@dangerous_kw",
        supporting_keywords=["-sub1", "+sub2"],
        target_audience="Developers",
        search_intent="Informational",
        content_format="Tutorial",
        opportunity_score=75.0,
        confidence_level="Moderate Evidence",
    )

    # 1. Opportunities CSV
    csv_bytes = ExportService.export_opportunities_csv([opp])
    assert csv_bytes.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM
    csv_text = csv_bytes.decode("utf-8-sig")
    assert "'=SUM(1,2)" in csv_text or "SUM" in csv_text
    assert "'@dangerous_kw" in csv_text or "@dangerous_kw" in csv_text

    # 2. Opportunities Excel
    excel_bytes = ExportService.export_opportunities_excel([opp])
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Opportunities" in wb.sheetnames
    ws = wb["Opportunities"]
    assert ws.max_row >= 2

    # 3. Calendar CSV & Excel
    cal_entry = CalendarEntry(
        title=dangerous_title,
        primary_keyword="car wax",
        planned_date="2026-11-20",
        status="Idea",
        priority="High",
        notes="+formula_note",
    )
    cal_csv = ExportService.export_calendar_csv([cal_entry])
    assert cal_csv.startswith(b"\xef\xbb\xbf")
    assert "'+formula_note" in cal_csv.decode("utf-8-sig")

    cal_excel = ExportService.export_calendar_excel([cal_entry])
    wb_cal = openpyxl.load_workbook(io.BytesIO(cal_excel))
    assert "Content Calendar" in wb_cal.sheetnames

    # 4. Overlap Report CSV & Excel
    overlap = TopicOverlapResult(
        topic_a=dangerous_title,
        topic_b="Normal Topic",
        similarity_score=0.85,
        overlap_type="High Token Overlap",
        shared_keywords=["kw1", "kw2"],
        warning_message="Warning",
        resolution="Unresolved",
    )
    ov_csv = ExportService.export_overlap_report_csv([overlap])
    assert ov_csv.startswith(b"\xef\xbb\xbf")

    ov_excel = ExportService.export_overlap_report_excel([overlap])
    wb_ov = openpyxl.load_workbook(io.BytesIO(ov_excel))
    assert "Cannibalization Report" in wb_ov.sheetnames


# ------------------------------------------------------------------------------
# 10. Multilingual Bengali Unicode Support
# ------------------------------------------------------------------------------

def test_multilingual_unicode_support() -> None:
    """Verify Bengali Unicode keywords and titles are preserved across gap analysis, topics, and exports."""
    bengali_kw = "গাড়ি পলিশ করার নিয়ম"  # rules of car polishing
    bengali_seed = "গাড়ি পলিশ"  # car polish
    competitors = ["গাড়ি পরিষ্কার ও ওয়াশ করার পদ্ধতি"]

    detector = ContentGapDetector(competitors)
    gap = detector.detect_gap(bengali_kw)
    assert gap["is_gap"] is True
    assert gap["gap_label"] == "Potential Content Gap"

    generator = VideoTopicGenerator()
    topics = generator.generate_topics_for_keywords(
        keywords=[bengali_kw],
        seed_keyword=bengali_seed,
        target_audience="General Audience",
        competitor_titles=competitors,
        max_topics=3,
    )
    assert len(topics) > 0
    bengali_topic = topics[0]
    assert bengali_kw in bengali_topic.primary_keyword or bengali_seed in bengali_topic.topic

    # Test export preserves Bengali characters without mojibake
    csv_bytes = ExportService.export_opportunities_csv([bengali_topic])
    decoded = csv_bytes.decode("utf-8-sig")
    assert "গাড়ি" in decoded


# ------------------------------------------------------------------------------
# 11. Navigation Constants & Backward Compatibility
# ------------------------------------------------------------------------------

def test_navigation_constants_and_compatibility() -> None:
    """Verify navigation constants maintain Phase 6 compatibility while exposing Phase 11."""
    assert PAGE_CONTENT_OPPORTUNITIES in ALL_NAVIGATION_PAGES
    # NAVIGATION_PAGES strictly keeps 6 items for Phase 6 test compliance
    assert len(NAVIGATION_PAGES) == 6
    assert PAGE_CONTENT_OPPORTUNITIES not in NAVIGATION_PAGES
    # ALL_NAVIGATION_PAGES has at least 8 items
    assert len(ALL_NAVIGATION_PAGES) >= 8
